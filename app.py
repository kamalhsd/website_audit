import fitz  # PyMuPDF for PDF to Image
from pdf2docx import Converter
from docx2pdf import convert as docx2pdf_convert
from htmldocx import HtmlToDocx
import docx
import tempfile
import pdfkit
import zipfile
import textstat
import whois
from sklearn.feature_extraction.text import CountVectorizer
from bs4 import BeautifulSoup
import requests
from urllib.parse import urlparse, urljoin
from concurrent.futures import ThreadPoolExecutor, as_completed
import time
from datetime import datetime
import re
import networkx as nx
import sys
import asyncio
import os

# --- FIX FOR WINDOWS & STREAMLIT COMPATIBILITY ---
if sys.platform.startswith("win"):
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
# --------------------------------------------------

import streamlit as st
import pandas as pd
from collections import deque, Counter
from concurrent.futures import TimeoutError
import json
from datetime import timedelta
import xml.etree.ElementTree as ET
from urllib.robotparser import RobotFileParser
import html
import threading
import sqlite3
import uuid
import traceback

# --- NEW IMPORTS FOR IMAGE OPTIMIZATION ---
try:
    from PIL import Image
    import io
    HAS_PIL = True
except ImportError:
    HAS_PIL = False
    
# --- AI & ML IMPORTS ---
try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    HAS_SKLEARN = True
except ImportError:
    HAS_SKLEARN = False

# --- SEMANTIC AI IMPORTS ---
try:
    from sentence_transformers import SentenceTransformer, util
    HAS_SENTENCE_TRANSFORMERS = True
except ImportError:
    HAS_SENTENCE_TRANSFORMERS = False

# --- GRAMMAR CHECKER IMPORT (Fixes NameError) ---
try:
    import language_tool_python
    HAS_LT = True
except ImportError:
    HAS_LT = False

# --- GSC IMPORTS ---
try:
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    HAS_GSC = True
except ImportError:
    HAS_GSC = False

# --- PLAYWRIGHT IMPORT (JS RENDERING) ---
try:
    from playwright.sync_api import sync_playwright
    HAS_PLAYWRIGHT = True
except ImportError:
    HAS_PLAYWRIGHT = False

# Page config
st.set_page_config(page_title="Battersea Crawler", layout="wide", page_icon="🐸")

# --- UI ENHANCEMENTS: THEME-AGNOSTIC CSS (DARK/LIGHT MODE READY) ---
st.markdown("""
<style>
/* Main Header Banner - Linear Gradients work in both modes */
.main-header {
    background: linear-gradient(135deg, #0f2027 0%, #203a43 50%, #2c5364 100%);
    padding: 30px;
    border-radius: 12px;
    text-align: center;
    box-shadow: 0 8px 16px rgba(0,0,0,0.15);
    margin-bottom: 30px;
}
.main-header h1 { color: #ffffff !important; margin: 0; font-size: 2.8rem; font-weight: 800; letter-spacing: 1px; padding-bottom: 5px; }
.main-header p { color: #4ade80 !important; font-size: 1.1rem; margin-top: 5px; font-weight: 600; text-transform: uppercase; letter-spacing: 2px;}

/* Metrics Cards (KPIs) - Using transparent RGBA so it adapts to Dark/Light mode */
div[data-testid="stMetric"] {
    background-color: rgba(128, 128, 128, 0.05);
    border: 1px solid rgba(128, 128, 128, 0.2);
    border-radius: 10px;
    padding: 15px 20px;
    box-shadow: 0 4px 6px rgba(0,0,0,0.04);
    transition: transform 0.2s ease, box-shadow 0.2s ease, background-color 0.2s ease;
}
div[data-testid="stMetric"]:hover {
    transform: translateY(-3px);
    box-shadow: 0 8px 15px rgba(0,0,0,0.1);
    background-color: rgba(128, 128, 128, 0.1);
}
/* Inherit text color for Metric values so it automatically flips to white/black */
div[data-testid="stMetricValue"] {
    font-weight: 700;
}

/* Tabs Styling - Theme Agnostic */
.stTabs [data-baseweb="tab-list"] {
    gap: 8px;
    padding: 5px 0px;
    background: transparent;
}
.stTabs [data-baseweb="tab"] {
    height: 45px;
    border-radius: 8px;
    background-color: rgba(128, 128, 128, 0.05) !important;
    border: 1px solid rgba(128, 128, 128, 0.2) !important;
    font-weight: 600;
    transition: all 0.2s ease;
    padding: 0 15px;
}
.stTabs [data-baseweb="tab"]:hover {
    background-color: rgba(128, 128, 128, 0.15) !important;
}
.stTabs [aria-selected="true"] {
    background: linear-gradient(135deg, #11998e 0%, #38ef7d 100%) !important;
    color: #ffffff !important;
    border: none !important;
    box-shadow: 0 4px 10px rgba(17, 153, 142, 0.3);
}

/* Dataframe container padding */
[data-testid="stDataFrame"] {
    border-radius: 8px;
    overflow: hidden;
    box-shadow: 0 2px 5px rgba(0,0,0,0.05);
}
</style>
""", unsafe_allow_html=True)

# Header
st.markdown("""
<div class="main-header">
    <h1>Battersea Crawler</h1>
    <p>Professional Edition • Full SEO Analysis</p>
</div>
""", unsafe_allow_html=True)


# --- GLOBAL CONFIG ---
write_lock = threading.Lock()

# Initialize session state
if 'db_file' not in st.session_state:
    st.session_state.db_file = f"battersea_data_{uuid.uuid4().hex}.db"

if 'crawl_data' not in st.session_state:
    st.session_state.crawl_data = []
if 'crawling' not in st.session_state:
    st.session_state.crawling = False
if 'stop_crawling' not in st.session_state:
    st.session_state.stop_crawling = False
if 'sitemap_urls_set' not in st.session_state:
    st.session_state.sitemap_urls_set = set()
if 'psi_results' not in st.session_state:
    st.session_state.psi_results = {}
if 'total_crawled_count' not in st.session_state:
    st.session_state.total_crawled_count = 0
if 'storage_mode' not in st.session_state:
    st.session_state.storage_mode = "RAM"
if 'img_size_cache' not in st.session_state:
    st.session_state.img_size_cache = {}
if 'img_real_dim_cache' not in st.session_state:
    st.session_state.img_real_dim_cache = {}
if 'img_rendered_cache' not in st.session_state:
    st.session_state.img_rendered_cache = {}
if 'img_status_cache' not in st.session_state:
    st.session_state.img_status_cache = {}
if 'link_status_cache' not in st.session_state:
    st.session_state.link_status_cache = {}
if 'gsc_service' not in st.session_state:
    st.session_state.gsc_service = None
if 'gsc_sites_list' not in st.session_state:
    st.session_state.gsc_sites_list = []
if 'gsc_merged_data' not in st.session_state:
    st.session_state.gsc_merged_data = pd.DataFrame()
if 'content_audit_data' not in st.session_state:
    st.session_state.content_audit_data = pd.DataFrame()
if 'grammar_audit_data' not in st.session_state:
    st.session_state.grammar_audit_data = pd.DataFrame()


# --- SQLITE HELPER FUNCTIONS ---
def init_db(db_path):
    conn = sqlite3.connect(db_path, check_same_thread=False)
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS pages (
            url TEXT PRIMARY KEY,
            original_url TEXT,
            final_url TEXT,
            status_code INTEGER,
            title TEXT,
            title_length INTEGER,
            meta_description TEXT,
            meta_desc_length INTEGER,
            canonical_url TEXT,
            meta_robots TEXT,
            is_noindex_follow BOOLEAN,
            is_noindex_nofollow BOOLEAN,
            h1_tags TEXT,
            h1_count INTEGER,
            h2_tags TEXT,
            h2_count INTEGER,
            h3_tags TEXT,
            h3_count INTEGER,
            h4_tags TEXT,
            h4_count INTEGER,
            word_count INTEGER,
            response_time REAL,
            content_length INTEGER,
            internal_links_count INTEGER,
            external_links_count INTEGER,
            internal_links TEXT,
            external_links TEXT,
            images TEXT,
            image_count INTEGER,
            images_without_alt INTEGER,
            schema_types TEXT,
            schema_dump TEXT,
            schema_count INTEGER,
            schema_validity TEXT,
            schema_errors TEXT,
            redirect_chain TEXT,
            redirect_count INTEGER,
            css_files INTEGER,
            js_files INTEGER,
            og_title TEXT,
            og_description TEXT,
            twitter_title TEXT,
            custom_extraction TEXT,
            custom_search_count INTEGER,
            indexability TEXT,
            crawl_timestamp TEXT,
            header_structure TEXT,
            scope_content TEXT,
            depth INTEGER, 
            error TEXT
        )
    ''')
    conn.commit()
    conn.close()

def save_row_to_db(data, db_path):
    row_data = data.copy()
    for key in ['internal_links', 'external_links', 'images', 'redirect_chain', 'schema_dump', 'header_structure']:
        if key in row_data:
            row_data[key] = json.dumps(row_data[key])
    
    if 'is_noindex_follow' in row_data:
        row_data['is_noindex_follow'] = 1 if row_data['is_noindex_follow'] else 0
    if 'is_noindex_nofollow' in row_data:
        row_data['is_noindex_nofollow'] = 1 if row_data['is_noindex_nofollow'] else 0

    columns = [
        'url', 'original_url', 'final_url', 'status_code', 'title', 'title_length', 
        'meta_description', 'meta_desc_length', 'canonical_url', 'meta_robots', 
        'is_noindex_follow', 'is_noindex_nofollow', 'h1_tags', 'h1_count', 
        'h2_tags', 'h2_count', 'h3_tags', 'h3_count', 'h4_tags', 'h4_count', 
        'word_count', 'response_time', 'content_length', 'internal_links_count', 
        'external_links_count', 'internal_links', 'external_links', 'images', 
        'image_count', 'images_without_alt', 'schema_types', 'schema_dump', 
        'schema_count', 'schema_validity', 'schema_errors', 'redirect_chain', 
        'redirect_count', 'css_files', 'js_files', 'og_title', 'og_description', 
        'twitter_title', 'custom_extraction', 'custom_search_count', 'indexability', 'crawl_timestamp', 
        'header_structure', 'scope_content', 'depth', 'error'
    ]
    
    filtered_data = {k: row_data.get(k) for k in columns}
    placeholders = ', '.join(['?'] * len(columns))
    col_names = ', '.join(columns)
    sql = f"INSERT OR REPLACE INTO pages ({col_names}) VALUES ({placeholders})"
    
    with write_lock:
        conn = sqlite3.connect(db_path, check_same_thread=False)
        try:
            conn.execute(sql, list(filtered_data.values()))
            conn.commit()
        except Exception as e:
            print(f"DB Error: {e}")
        finally:
            conn.close()

# --- GSC HELPER FUNCTIONS ---
def authenticate_gsc(json_content):
    try:
        info = json.loads(json_content)
        creds = service_account.Credentials.from_service_account_info(
            info, scopes=['https://www.googleapis.com/auth/webmasters.readonly']
        )
        service = build('searchconsole', 'v1', credentials=creds)
        return service
    except Exception as e:
        st.error(f"Authentication Error: {e}")
        return None

def list_gsc_sites(service):
    try:
        site_list = service.sites().list().execute()
        return [s['siteUrl'] for s in site_list.get('siteEntry', [])]
    except Exception as e:
        st.error(f"Error listing sites: {e}")
        return []

def fetch_gsc_data(service, site_url, start_date, end_date):
    try:
        s_date = start_date.strftime('%Y-%m-%d')
        e_date = end_date.strftime('%Y-%m-%d')
    except AttributeError:
        s_date = str(start_date)
        e_date = str(end_date)

    request = {
        'startDate': s_date, 'endDate': e_date,
        'dimensions': ['page'], 'rowLimit': 25000
    }
    
    try:
        response = service.searchanalytics().query(siteUrl=site_url, body=request).execute()
        rows = response.get('rows', [])
        data = []
        for row in rows:
            data.append({
                'url': row['keys'][0],
                'GSC Clicks': row['clicks'],
                'GSC Impressions': row['impressions'],
                'GSC CTR': row['ctr'],
                'GSC Position': row['position']
            })
        return pd.DataFrame(data)
    except Exception as e:
        st.error(f"Error fetching data: {e}")
        return pd.DataFrame()

def inspect_url_indexing(service, site_url, url_list):
    results = []
    progress_bar = st.progress(0)
    status_text = st.empty()
    total = len(url_list)
    
    for i, u in enumerate(url_list):
        progress_bar.progress((i + 1) / total)
        status_text.text(f"Inspecting {i+1}/{total}: {u}")
        try:
            req = {'inspectionUrl': u, 'siteUrl': site_url}
            resp = service.urlInspection().index().inspect(body=req).execute()
            inspection_res = resp.get('inspectionResult', {})
            index_res = inspection_res.get('indexStatusResult', {})
            mobile_res = inspection_res.get('mobileUsabilityResult', {})
            results.append({
                'url': u,
                'Coverage State': index_res.get('coverageState', 'Unknown'),
                'Indexing Status': index_res.get('indexingState', 'Unknown'),
                'Google Canonical': index_res.get('googleCanonical', 'N/A'),
                'User Canonical': index_res.get('userCanonical', 'N/A'),
                'Last Crawl Time': index_res.get('lastCrawlTime', 'Never'),
                'Mobile Usability': mobile_res.get('verdict', 'Not Tested')
            })
            time.sleep(0.2) 
        except Exception as e:
            err_msg = str(e)
            if "403" in err_msg: err_msg = "Permission Denied"
            elif "429" in err_msg: err_msg = "Quota Exceeded"
            results.append({'url': u, 'Coverage State': f"Error: {err_msg}"})
            
    status_text.success("✅ Inspection Complete!")
    return pd.DataFrame(results)
    
def submit_to_indexing_api(json_content, url_list, action="URL_UPDATED"):
    """Submits URLs directly to the Google Indexing API."""
    results = []
    try:
        info = json.loads(json_content)
        # Note: The Indexing API requires a specific scope, different from GSC!
        creds = service_account.Credentials.from_service_account_info(
            info, scopes=['https://www.googleapis.com/auth/indexing']
        )
        service = build('indexing', 'v3', credentials=creds)
        
        for url in url_list:
            try:
                body = {"url": url, "type": action}
                response = service.urlNotifications().publish(body=body).execute()
                results.append({"URL": url, "Status": "✅ Success", "Details": "Submitted for Crawl"})
            except Exception as e:
                # Catch API errors (like 429 Quota Exceeded or 403 Permission Denied)
                err_msg = str(e)
                if "429" in err_msg: err_msg = "Quota Exceeded (Daily Limit Reached)"
                elif "403" in err_msg: err_msg = "Permission Denied (Ensure Service Account is added as Owner)"
                results.append({"URL": url, "Status": "❌ Failed", "Details": err_msg})
            
            # Gentle delay to prevent overwhelming the API
            time.sleep(0.5) 
            
    except Exception as e:
        return pd.DataFrame([{"URL": "Authentication Error", "Status": "❌ Failed", "Details": str(e)}])
        
    return pd.DataFrame(results)

# --- AI HELPER FUNCTIONS ---
@st.cache_resource
def load_embedding_model():
    return SentenceTransformer('all-MiniLM-L6-v2')

def generate_interlink_suggestions(df, min_score=40, max_suggestions=10):
    if df.empty: return pd.DataFrame()
    
    if not HAS_SENTENCE_TRANSFORMERS:
        st.error("❌ 'sentence-transformers' library missing. Run `pip install sentence-transformers`")
        return pd.DataFrame()

    df['target_topic'] = df['title'].fillna('') + " " + df['h1_tags'].fillna('')
    df['source_context'] = df['scope_content'].fillna('')
    
    mask = df['source_context'] == ''
    df.loc[mask, 'source_context'] = df.loc[mask, 'title'].fillna('') + " " + df.loc[mask, 'meta_description'].fillna('')

    valid_targets = df[df['target_topic'].str.strip() != '']
    valid_sources = df[df['source_context'].str.strip() != '']
    if valid_targets.empty or valid_sources.empty: return pd.DataFrame()

    model = load_embedding_model()

    target_texts = df['target_topic'].tolist()
    source_texts = df['source_context'].str[:3000].tolist()
    
    target_embeddings = model.encode(target_texts, convert_to_tensor=True)
    source_embeddings = model.encode(source_texts, convert_to_tensor=True)

    similarity_matrix = util.cos_sim(source_embeddings, target_embeddings).cpu().numpy()

    suggestions = []
    existing_links = set()
    
    for _, row in df.iterrows():
        links = row.get('internal_links', [])
        if isinstance(links, str):
            try: links = json.loads(links)
            except: links = []
        for link in links:
            if isinstance(link, dict) and 'url' in link:
                existing_links.add((row['url'], link['url']))

    for idx, row in df.iterrows():
        source_url = row['url']
        scores = similarity_matrix[idx]
        
        top_indices = scores.argsort()[::-1][:50] 
        count = 0
        
        for target_idx in top_indices:
            if count >= max_suggestions: break
            
            target_url = df.iloc[target_idx]['url']
            score = round(float(scores[target_idx]) * 100, 1) 
            
            if source_url == target_url: continue
            if score < min_score: continue
            if (source_url, target_url) in existing_links: continue
            
            suggestions.append({
                'Source URL': source_url,
                'Source Content Preview': row['source_context'][:60] + "...",
                'Suggested Target URL': target_url,
                'Target Topic': df.iloc[target_idx]['target_topic'][:60] + "...",
                'Relevance Score': score
            })
            count += 1
            
    return pd.DataFrame(suggestions)

def analyze_content_cannibalization(df, merge_threshold=0.50, duplicate_threshold=0.85):
    if df.empty: return pd.DataFrame()
    
    valid_df = df[df['scope_content'].str.len() > 100].copy().reset_index(drop=True)
    if len(valid_df) < 2: return pd.DataFrame()

    valid_df['analysis_text'] = (
        (valid_df['title'].fillna('') + " ") * 3 + 
        (valid_df['h1_tags'].fillna('') + " ") * 3 + 
        valid_df['scope_content'].fillna('').str[:5000]
    )

    tfidf = TfidfVectorizer(stop_words='english', min_df=1)
    try:
        tfidf_matrix = tfidf.fit_transform(valid_df['analysis_text'])
    except: return pd.DataFrame()

    cosine_sim = cosine_similarity(tfidf_matrix, tfidf_matrix)
    results = []
    num_rows = len(valid_df)
    
    for i in range(num_rows):
        for j in range(i + 1, num_rows):
            score = cosine_sim[i, j]
            if score < merge_threshold: continue
            
            if score >= duplicate_threshold:
                action = "🚨 Remove/Redirect (Duplicate)"
                reason = "Content is nearly identical (>{}%)".format(int(duplicate_threshold*100))
            else:
                action = "🤝 Merge Content"
                reason = "Topics overlap significantly (>{}%)".format(int(merge_threshold*100))

            results.append({
                'Page A': valid_df.iloc[i]['url'],
                'Page A Title': valid_df.iloc[i]['title'],
                'Page B': valid_df.iloc[j]['url'],
                'Page B Title': valid_df.iloc[j]['title'],
                'Similarity': round(score * 100, 1),
                'Recommendation': action,
                'Reason': reason
            })
            
    return pd.DataFrame(results).sort_values(by='Similarity', ascending=False)

def generate_ai_meta(provider, api_key, model_name, endpoint_url, prompt_text, system_instruction="You are an SEO expert."):
    try:
        if provider == "Gemini":
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={api_key}"
            headers = {"Content-Type": "application/json"}
            payload = {"contents": [{"parts": [{"text": f"{system_instruction}\n\n{prompt_text}"}]}]}
            response = requests.post(url, headers=headers, json=payload, timeout=30)
            if response.status_code == 200:
                data = response.json()
                return data['candidates'][0]['content']['parts'][0]['text'].strip()
            else:
                return f"Error {response.status_code}: {response.text}"
        elif provider == "OpenAI Compatible (Groq/Ollama/OpenAI)":
            headers = {"Content-Type": "application/json", "Authorization": f"Bearer {api_key}" if api_key else ""}
            payload = {
                "model": model_name,
                "messages": [{"role": "system", "content": system_instruction}, {"role": "user", "content": prompt_text}],
                "temperature": 0.7
            }
            response = requests.post(endpoint_url, headers=headers, json=payload, timeout=30)
            if response.status_code == 200:
                data = response.json()
                return data['choices'][0]['message']['content'].strip()
            else:
                return f"Error {response.status_code}: {response.text}"
    except Exception as e:
        return f"Exception: {str(e)}"
    return "Unknown Error"

def analyze_content_freshness(url, title, content, provider, api_key, model_name, endpoint_url):
    current_date = datetime.now().strftime("%B %Y")
    safe_content = str(content) if pd.notna(content) else ""
    
    prompt = f"""
    You are an expert SEO Content Strategist. 
    Today's Date: {current_date}
    Analyze the following page content and determine its relevance for TODAY's users.
    URL: {url}
    Title: {title}
    Content Snippet: {safe_content[:1500]}
    
    Decision Criteria:
    - KEEP: Content is evergreen, still accurate, or historically valuable.
    - UPDATE: Content is good but contains old dates, old prices, or missing newer facts.
    - REMOVE: Content is completely outdated.
    
    Provide your response in this EXACT format. Do not use any introductory text:
    DECISION: [KEEP/UPDATE/REMOVE]
    REASON: [Short explanation why]
    ACTION: [What exactly to add or change, or why to delete]
    """
    
    try:
        response = generate_ai_meta(provider, api_key, model_name, endpoint_url, prompt, "You are a content auditor.")
        
        if response.startswith("Error") or response.startswith("Exception"):
            return {"URL": url, "Decision": "Error", "Reason": response, "Action Suggestion": "Check API Key/Limits", "Raw AI Output": response}

        res_dict = {
            "URL": url, 
            "Decision": "UNKNOWN", 
            "Reason": "Format mismatched. See Raw Output.", 
            "Action Suggestion": "N/A",
            "Raw AI Output": response 
        }
        
        lines = response.split('\n')
        for line in lines:
            clean_line = line.replace('*', '').strip()
            upper_line = clean_line.upper()
            
            if upper_line.startswith("DECISION:") or upper_line.startswith("DECISION :"): 
                res_dict["Decision"] = clean_line.split(":", 1)[1].strip().upper()
                if "Format mismatched" in res_dict["Reason"]:
                    res_dict["Reason"] = ""
            elif upper_line.startswith("REASON:") or upper_line.startswith("REASON :"): 
                res_dict["Reason"] = clean_line.split(":", 1)[1].strip()
            elif upper_line.startswith("ACTION:") or upper_line.startswith("ACTION :"): 
                res_dict["Action Suggestion"] = clean_line.split(":", 1)[1].strip()
                
        if "KEEP" in res_dict["Decision"]: res_dict["Decision"] = "KEEP"
        elif "UPDATE" in res_dict["Decision"]: res_dict["Decision"] = "UPDATE"
        elif "REMOVE" in res_dict["Decision"]: res_dict["Decision"] = "REMOVE"
                
        return res_dict
        
    except Exception as e:
        return {"URL": url, "Decision": "Error", "Reason": f"Crash: {str(e)}", "Action Suggestion": "N/A", "Raw AI Output": "N/A"}

# --- GRAMMAR FALLBACK HELPER (UPDATED) ---
def check_grammar_cloud(text):
    url = "https://api.languagetool.org/v2/check"
    data = {
        'text': text,
        'language': 'en-US'
    }
    try:
        response = requests.post(url, data=data, timeout=5)
        if response.status_code == 200:
            result = response.json()
            matches = []
            for m in result.get('matches', []):
                replacements = [r['value'] for r in m.get('replacements', [])][:3]
                matches.append({
                    'message': m.get('message', ''),
                    'replacements': replacements,
                    'offset': m.get('offset', 0),
                    'length': m.get('length', 0)
                })
            return matches
        return []
    except Exception:
        return []
        
def calculate_internal_pagerank(df):
    try:
        G = nx.DiGraph()
        
        for _, row in df.iterrows():
            source = row['url']
            G.add_node(source)
            
            links = row.get('internal_links', [])
            if isinstance(links, str):
                try: links = json.loads(links)
                except: links = []
            elif isinstance(links, list):
                pass
            else:
                links = []
                
            for link in links:
                if isinstance(link, dict):
                    target = link.get('url')
                    if target:
                        G.add_edge(source, target)
        
        try:
            pagerank_scores = nx.pagerank(G, alpha=0.85, max_iter=100)
        except:
            return {}
        
        if not pagerank_scores: return {}
        
        max_score = max(pagerank_scores.values())
        min_score = min(pagerank_scores.values())
        
        if max_score == min_score: return {k: 50 for k in pagerank_scores}

        normalized_scores = {
            k: int(((v - min_score) / (max_score - min_score)) * 100) 
            for k, v in pagerank_scores.items()
        }
        
        return normalized_scores
    except Exception as e:
        st.error(f"PageRank Error: {e}")
        return {}
        
def get_domain_age(domain):
    try:
        w = whois.whois(domain)
        creation_date = w.creation_date
        if isinstance(creation_date, list): creation_date = creation_date[0]
        if creation_date:
            age = (datetime.now() - creation_date).days / 365.25
            return round(age, 1)
        return "Unknown"
    except: return "Hidden"

def extract_publish_date(soup):
    meta_tags = soup.find_all('meta')
    for tag in meta_tags:
        prop = tag.get('property', '').lower()
        name = tag.get('name', '').lower()
        if 'published_time' in prop or 'modified_time' in prop or 'date' in name:
            content = tag.get('content', '')
            if content: return content[:10]
    return "Not Found"

def count_fuzzy_match(text, keyword):
    sentences = text.split('.')
    kw_words = [w.lower() for w in keyword.split()]
    count = 0
    for sentence in sentences:
        if all(w in sentence.lower() for w in kw_words):
            count += 1
    return count

def deep_crawl_for_inlinks(target_url, max_pages=300):
    parsed_target = urlparse(target_url)
    domain = f"{parsed_target.scheme}://{parsed_target.netloc}"
    target_path = parsed_target.path
    
    visited = set()
    queue = [domain]
    inlink_count = 0
    
    session = requests.Session()
    session.headers.update({'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
    
    while queue and len(visited) < max_pages:
        current = queue.pop(0)
        if current in visited: continue
        visited.add(current)
        
        try:
            res = session.get(current, timeout=5)
            if res.status_code == 200:
                soup = BeautifulSoup(res.text, 'html.parser')
                links = [a.get('href', '') for a in soup.find_all('a', href=True)]
                
                for link in links:
                    if target_url in link or (target_path in link and len(target_path) > 1):
                        if current != target_url:
                            inlink_count += 1
                        break
                        
                for link in links:
                    full_url = urljoin(domain, link)
                    if full_url.startswith(domain) and full_url not in visited and full_url not in queue:
                        queue.append(full_url)
        except: pass
        
    return inlink_count
    
# --- CRAWLER CLASS ---
class UltraFrogCrawler:
    def __init__(self, max_urls=100000, ignore_robots=False, crawl_scope="subfolder", custom_selector=None, link_selector=None, search_config=None, use_js=False):
        self.max_urls = max_urls
        self.ignore_robots = ignore_robots
        self.crawl_scope = crawl_scope
        self.custom_selector = custom_selector
        self.link_selector = link_selector
        self.search_config = search_config
        self.use_js = use_js 
        
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Linux; Android 10; K) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Mobile Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
            'Accept-Language': 'en-US,en;q=0.5',
            'Accept-Encoding': 'gzip, deflate',
            'Connection': 'keep-alive',
            'Upgrade-Insecure-Requests': '1'
        })
        adapter = requests.adapters.HTTPAdapter(pool_connections=20, pool_maxsize=20, max_retries=1)
        self.session.mount('http://', adapter)
        self.session.mount('https://', adapter)
        self.robots_cache = {}
        self.base_domain = None
        self.base_path = None
    
    def set_base_url(self, url):
        parsed = urlparse(url)
        self.base_domain = parsed.netloc
        self.base_path = parsed.path.rstrip('/')
    
    def should_crawl_url(self, url):
        parsed = urlparse(url)
        base_clean = self.base_domain.replace('www.', '')
        target_clean = parsed.netloc.replace('www.', '')
        if self.crawl_scope == "exact":
            return url == urljoin(f"https://{self.base_domain}", self.base_path)
        elif self.crawl_scope == "subdomain":
            return base_clean in target_clean
        else:
            return (base_clean == target_clean and parsed.path.startswith(self.base_path))
    
    def can_fetch(self, url):
        if self.ignore_robots: return True
        try:
            domain = urlparse(url).netloc
            if domain not in self.robots_cache:
                try:
                    rp = RobotFileParser()
                    rp.set_url(f"https://{domain}/robots.txt")
                    rp.read()
                    self.robots_cache[domain] = rp
                except: self.robots_cache[domain] = None
            if self.robots_cache[domain] is None: return True
            return self.robots_cache[domain].can_fetch('*', url)
        except: return True

    def smart_clean(self, text):
        if not text: return ""
        text = str(text)
        text = html.unescape(text)
        text = re.sub(r'[\r\n\t]+', ' ', text)
        text = re.sub(r'\s+', ' ', text)
        return text.strip()

    def get_css_path(self, element):
        path = []
        for parent in element.parents:
            if parent.name == '[document]': break
            selector = parent.name
            if parent.get('id'): selector += f"#{parent['id']}"
            elif parent.get('class'): selector += f".{parent['class'][0]}"
            path.append(selector)
        path.reverse()
        element_selector = element.name
        if element.get('class'): element_selector += f".{element['class'][0]}"
        path.append(element_selector)
        return " > ".join(path)

    def extract_sitemap_urls(self, sitemap_url):
        urls = []
        try:
            response = self.session.get(sitemap_url, timeout=8)
            if response.status_code == 200:
                root = ET.fromstring(response.content)
                namespaces = {'sitemap': 'http://www.sitemaps.org/schemas/sitemap/0.9'}
                sitemapindex = root.findall('.//sitemap:sitemap', namespaces)
                if sitemapindex:
                    for sitemap in sitemapindex:
                        loc = sitemap.find('sitemap:loc', namespaces)
                        if loc is not None and loc.text: urls.extend(self.extract_sitemap_urls(loc.text.strip()))
                else:
                    url_elements = root.findall('.//sitemap:url', namespaces)
                    for url_elem in url_elements:
                        loc = url_elem.find('sitemap:loc', namespaces)
                        if loc is not None and loc.text: urls.append(loc.text.strip())
                        
        except Exception as e: st.error(f"Error parsing sitemap: {e}")
        return urls
        
    def get_file_size(self, url):
        try:
            r = self.session.head(url, timeout=2)
            if 'content-length' in r.headers: return round(int(r.headers['content-length']) / 1024, 2)
        except: pass
        return 0

    def fetch_with_playwright(self, url):
        if not HAS_PLAYWRIGHT:
            return None, "Playwright library not found."
        
        try:
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                context = browser.new_context(user_agent='Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36')
                page = context.new_page()
                page.goto(url, wait_until="networkidle", timeout=30000) 
                content = page.content() 
                
                class MockResponse:
                    def __init__(self, url, content, status_code):
                        self.url = url
                        self.content = content.encode('utf-8')
                        self.text = content
                        self.status_code = status_code
                        self.history = []
                        self.elapsed = type('obj', (object,), {'total_seconds': lambda: 1.0})
                
                browser.close()
                return MockResponse(url, content, 200), None
        except Exception as e:
            return None, str(e)

    def extract_page_data(self, url):
        try:
            if self.use_js:
                response, error_msg = self.fetch_with_playwright(url)
                if not response:
                    raise Exception(error_msg)
            else:
                response = self.session.get(url, timeout=10, allow_redirects=True)
            
            actual_status_code = response.status_code
            if hasattr(response, 'history') and response.history:
                actual_status_code = response.history[0].status_code

            soup = BeautifulSoup(response.content, 'html.parser')
            
            custom_search_count = 0
            if self.search_config and self.search_config.get('query'):
                query = self.search_config['query']
                mode = self.search_config['mode']
                scope = self.search_config['scope']
                
                search_target = ""
                if mode == "CSS Selector Existence":
                    try:
                        elements = soup.select(query)
                        custom_search_count = len(elements)
                    except: custom_search_count = 0
                else:
                    if scope == "HTML Source Code": search_target = response.text
                    else: search_target = soup.get_text()
                    
                    if mode == "Text (Case Insensitive)": custom_search_count = search_target.lower().count(query.lower())
                    elif mode == "Text (Case Sensitive)": custom_search_count = search_target.count(query)
                    elif mode == "Regex":
                        try: custom_search_count = len(re.findall(query, search_target))
                        except: custom_search_count = 0

            scope_content = ""
            if self.link_selector:
                scoped_element = soup.select_one(self.link_selector)
                if scoped_element: scope_content = self.smart_clean(scoped_element.get_text())
            else:
                if soup.body: scope_content = self.smart_clean(soup.body.get_text())[:50000]

            header_structure = []
            all_headers = soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6'])
            for tag in all_headers:
                header_structure.append({
                    'tag': tag.name,
                    'level': int(tag.name[1]), 
                    'text': self.smart_clean(tag.get_text())[:100]
                })
            
            title = soup.find('title')
            title_text = self.smart_clean(title.get_text()) if title else ""
            
            meta_desc = soup.find('meta', attrs={'name': 'description'})
            meta_desc_text = self.smart_clean(meta_desc.get('content', '')) if meta_desc else ""
            
            canonical = soup.find('link', attrs={'rel': 'canonical'})
            canonical_url = canonical.get('href') if canonical else ""
            
            meta_robots = soup.find('meta', attrs={'name': 'robots'})
            robots_content = meta_robots.get('content', '') if meta_robots else ""
            
            is_noindex_follow = False
            is_noindex_nofollow = False
            if robots_content:
                directives = [d.strip().lower() for d in robots_content.split(',')]
                if 'noindex' in directives and 'follow' in directives: is_noindex_follow = True
                elif 'noindex' in directives and 'nofollow' in directives: is_noindex_nofollow = True

            h1_tags = [self.smart_clean(h1.get_text()) for h1 in soup.find_all('h1')]
            h2_tags = [self.smart_clean(h2.get_text()) for h2 in soup.find_all('h2')]
            h3_tags = [self.smart_clean(h3.get_text()) for h3 in soup.find_all('h3')]
            h4_tags = [self.smart_clean(h4.get_text()) for h4 in soup.find_all('h4')]
            
            custom_data_dict = {}
            if self.custom_selector and isinstance(self.custom_selector, list):
                for rule in self.custom_selector:
                    try:
                        elements = soup.select(rule['selector'])
                        values = []
                        for el in elements:
                            val = ""
                            if rule['type'] == "Text Content":
                                val = self.smart_clean(el.get_text())
                            elif rule['type'] == "Inner HTML":
                                val = str(el.encode_contents().decode('utf-8')).strip()
                            elif rule['type'] == "HTML Element":
                                val = str(el).strip()
                            elif rule['type'] == "Attribute Value":
                                val = el.get(rule['attribute'], '')
                            
                            if val: values.append(str(val))
                        
                        custom_data_dict[rule['name']] = "; ".join(values) if values else ""
                    except Exception:
                        custom_data_dict[rule['name']] = "Error"

            custom_extraction_json = json.dumps(custom_data_dict)

            internal_links = []
            external_links = []
            base_domain_clean = urlparse(response.url).netloc.replace('www.', '')
            
            search_area = soup
            if self.link_selector:
                specific_section = soup.select_one(self.link_selector)
                if specific_section: search_area = specific_section
            
            for link in search_area.find_all('a', href=True):
                # FIX: Strip whitespace from the raw href before and after processing
                raw_href = link['href'].strip()
                href = urljoin(response.url, raw_href) 
                href = href.split('#')[0].strip()
                if not href: continue

                link_text = self.smart_clean(link.get_text())[:100]
                css_path = self.get_css_path(link)
                
                rel_attr = link.get('rel', [])
                target_attr = link.get('target', '')
                rel_str = ' '.join(rel_attr) if isinstance(rel_attr, list) else str(rel_attr)
                
                link_rel_status = "dofollow"
                if 'nofollow' in rel_str.lower(): link_rel_status = "nofollow"
                if 'sponsored' in rel_str.lower(): link_rel_status += ", sponsored"
                if 'ugc' in rel_str.lower(): link_rel_status += ", ugc"
                
                link_target = "_blank" if target_attr == "_blank" else "_self"
                
                placement = "Body"
                path_lower = css_path.lower()
                if "footer" in path_lower: placement = "Footer"
                elif "header" in path_lower or "nav" in path_lower or "menu" in path_lower: placement = "Header"
                elif "sidebar" in path_lower or "aside" in path_lower: placement = "Sidebar"

                link_data = {
                    'url': href, 'anchor_text': link_text, 'css_path': css_path,
                    'placement': placement, 'rel_status': link_rel_status, 'target': link_target
                }
                
                link_netloc = urlparse(href).netloc.replace('www.', '')
                if link_netloc == base_domain_clean: internal_links.append(link_data)
                else: external_links.append(link_data)
            
            images = []
            for img in soup.find_all('img'):
                img_src = urljoin(response.url, img.get('src', ''))
                images.append({
                    'src': img_src,
                    'alt': self.smart_clean(img.get('alt', '')),
                    'title': self.smart_clean(img.get('title', '')),
                    'width': img.get('width', ''),
                    'height': img.get('height', ''),
                    'size_kb': 0
                })
            
            scripts = soup.find_all('script', type='application/ld+json')
            schema_types = []
            schema_dump = [] 
            schema_validity = "No Schema"
            schema_errors = []
            
            if scripts:
                schema_validity = "Valid"
                for script in scripts:
                    try:
                        if script.string:
                            data = json.loads(script.string)
                            schema_dump.append(data)
                            
                            if isinstance(data, list):
                                for item in data:
                                    if isinstance(item, dict) and '@type' in item:
                                        t = item['@type']
                                        schema_types.extend(t if isinstance(t, list) else [t])
                            
                            elif isinstance(data, dict) and '@graph' in data and isinstance(data['@graph'], list):
                                for item in data['@graph']:
                                    if isinstance(item, dict) and '@type' in item:
                                        t = item['@type']
                                        schema_types.extend(t if isinstance(t, list) else [t])
                                        
                            elif isinstance(data, dict) and '@type' in data:
                                t = data['@type']
                                schema_types.extend(t if isinstance(t, list) else [t])
                                
                    except json.JSONDecodeError as e:
                        schema_validity = "Invalid JSON"
                        schema_errors.append(str(e))
                    except Exception as e:
                        schema_validity = "Error"
                        schema_errors.append(str(e))

            css_files = len(soup.find_all('link', attrs={'rel': 'stylesheet'}))
            js_files = len(soup.find_all('script', src=True))
            text_content = soup.get_text()
            word_count = len(text_content.split())
            
            redirect_chain = []
            if hasattr(response, 'history') and response.history:
                for i, resp in enumerate(response.history):
                    redirect_chain.append({
                        'step': i + 1, 'from_url': resp.url,
                        'status_code': resp.status_code,
                        'redirect_type': '301 Permanent' if resp.status_code == 301 else f'{resp.status_code}'
                    })
            
            return {
                'url': url, 
                'original_url': url,
                'final_url': response.url,
                'status_code': actual_status_code,
                'title': title_text, 'title_length': len(title_text),
                'meta_description': meta_desc_text, 'meta_desc_length': len(meta_desc_text),
                'canonical_url': canonical_url, 'meta_robots': robots_content,
                'is_noindex_follow': is_noindex_follow, 'is_noindex_nofollow': is_noindex_nofollow,
                'h1_tags': '; '.join(h1_tags), 'h1_count': len(h1_tags),
                'h2_tags': '; '.join(h2_tags), 'h2_count': len(h2_tags),
                'h3_tags': '; '.join(h3_tags), 'h3_count': len(h3_tags),
                'h4_tags': '; '.join(h4_tags), 'h4_count': len(h4_tags),
                'word_count': word_count, 'response_time': 1.0, 
                'content_length': len(response.content) if hasattr(response, 'content') else len(response.text),
                'internal_links_count': len(internal_links), 'external_links_count': len(external_links),
                'internal_links': internal_links, 'external_links': external_links,
                'images': images, 'image_count': len(images),
                'images_without_alt': len([img for img in images if not img['alt']]),
                'schema_types': '; '.join(list(set(schema_types))), 'schema_dump': schema_dump,
                'schema_count': len(schema_types), 'schema_validity': schema_validity,
                'schema_errors': '; '.join(schema_errors), 'redirect_chain': redirect_chain,
                'redirect_count': len(redirect_chain), 'css_files': css_files, 'js_files': js_files,
                'og_title': soup.find('meta', attrs={'property': 'og:title'}).get('content', '') if soup.find('meta', attrs={'property': 'og:title'}) else '',
                'og_description': soup.find('meta', attrs={'property': 'og:description'}).get('content', '') if soup.find('meta', attrs={'property': 'og:description'}) else '',
                'twitter_title': soup.find('meta', attrs={'name': 'twitter:title'}).get('content', '') if soup.find('meta', attrs={'name': 'twitter:title'}) else '',
                'custom_extraction': custom_extraction_json, 
                'custom_search_count': custom_search_count,
                'indexability': self.get_indexability_status(actual_status_code, robots_content),
                'crawl_timestamp': datetime.now().isoformat(), 
                'header_structure': header_structure, 
                'scope_content': scope_content, 
                'error': ''
            }
        except Exception as e:
            return {
                'url': url, 'original_url': url, 'final_url': url, 'status_code': 0, 'error': str(e),
                'title': '', 'title_length': 0, 'meta_description': '', 'meta_desc_length': 0,
                'canonical_url': '', 'meta_robots': '', 'is_noindex_follow': False, 'is_noindex_nofollow': False,
                'h1_tags': '', 'h1_count': 0, 'h2_tags': '', 'h2_count': 0, 'h3_tags': '', 'h3_count': 0,
                'h4_tags': '', 'h4_count': 0, 'word_count': 0, 'response_time': 0,
                'content_length': 0, 'internal_links_count': 0, 'external_links_count': 0,
                'internal_links': [], 'external_links': [], 'images': [], 'image_count': 0,
                'images_without_alt': 0, 'schema_types': '', 'schema_dump': [], 'schema_count': 0, 
                'schema_validity': 'Error', 'schema_errors': '',
                'redirect_chain': [], 'redirect_count': 0, 'css_files': 0, 'js_files': 0,
                'og_title': '', 'og_description': '', 'twitter_title': '',
                'custom_extraction': '', 'custom_search_count': 0, 'indexability': 'Error', 'crawl_timestamp': datetime.now().isoformat(),
                'header_structure': [], 'scope_content': ''
            }
    
    def get_indexability_status(self, status_code, robots_content):
        if status_code != 200:
            return 'Non-Indexable'
        if 'noindex' in robots_content.lower():
            return 'Non-Indexable'
        return 'Indexable'

# --- PSI HELPER FUNCTION ---
def run_psi_test(url, api_key, strategy="mobile"):
    if not api_key: return {"error": "No API Key Provided"}
    api_url = f"https://www.googleapis.com/pagespeedonline/v5/runPagespeed?url={url}&strategy={strategy}&key={api_key}"
    try:
        response = requests.get(api_url)
        data = response.json()
        if "error" in data: 
            return {"error": data["error"]["message"]}
        lh_result = data["lighthouseResult"]
        return {
            "Score": int(lh_result["categories"]["performance"]["score"] * 100),
            "LCP": lh_result["audits"]["largest-contentful-paint"]["displayValue"],
            "FCP": lh_result["audits"]["first-contentful-paint"]["displayValue"],
            "CLS": lh_result["audits"]["cumulative-layout-shift"]["displayValue"],
            "INP": lh_result["audits"].get("interaction-to-next-paint", {}).get("displayValue", "N/A"),
            "TTI": lh_result["audits"]["interactive"]["displayValue"]
        }
    except Exception as e: 
        return {"error": str(e)}

# --- HELPER FOR HEADER ANALYSIS ---
def analyze_heading_structure(structure):
    issues = []
    h1_count = 0
    prev_level = 0
    
    if not structure:
        return ["⚠️ No Heading Tags Found"], 0

    for h in structure:
        curr_level = h['level']
        if curr_level == 1:
            h1_count += 1
        if curr_level - prev_level > 1 and prev_level != 0:
            issues.append(f"⚠️ Skipped Level: {h['tag'].upper()} follows H{prev_level} (Should be H{prev_level+1})")
        prev_level = curr_level

    if h1_count == 0:
        issues.insert(0, "❌ Missing H1 Tag")
    elif h1_count > 1:
        issues.insert(0, f"❌ Multiple H1 Tags Found ({h1_count})")
        
    return issues, h1_count

# --- PRODUCTION HELPER: VISUAL DIMENSIONS (PLAYWRIGHT) ---
def measure_rendered_images(url_list, progress_callback=None):
    if not HAS_PLAYWRIGHT:
        return {}, "Playwright not installed"

    results = {}
    total_images_found = 0
    
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            try:
                browser = p.chromium.launch(
                    headless=True,
                    args=["--disable-blink-features=AutomationControlled"]
                )
            except Exception as e:
                return {}, f"BROWSER LAUNCH FAILED: {str(e)}\n\nTry running: playwright install"
            
            REAL_USER_AGENT = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
            
            context_desk = browser.new_context(
                viewport={"width": 1920, "height": 1080},
                user_agent=REAL_USER_AGENT,
                locale='en-US'
            )
            page_desk = context_desk.new_page()

            context_mob = browser.new_context(
                viewport={"width": 390, "height": 844},
                user_agent="Mozilla/5.0 (iPhone; CPU iPhone OS 16_0 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/16.0 Mobile/15E148 Safari/604.1",
                is_mobile=True,
                locale='en-US'
            )
            page_mob = context_mob.new_page()

            total = len(url_list)
            for i, url in enumerate(url_list):
                if progress_callback:
                    progress_callback(i, total, url)

                try:
                    page_desk.goto(url, wait_until="domcontentloaded", timeout=25000)
                    
                    page_desk.evaluate("""
                        async () => {
                            const delay = ms => new Promise(resolve => setTimeout(resolve, ms));
                            for (let i = 0; i < document.body.scrollHeight; i += 500) {
                                window.scrollTo(0, i);
                                await delay(50); 
                            }
                        }
                    """)
                    time.sleep(0.5) 
                    
                    imgs_desk = page_desk.evaluate("""
                        Array.from(document.querySelectorAll('img')).map(img => {
                            const rect = img.getBoundingClientRect();
                            return {
                                src: img.currentSrc || img.src, 
                                width: Math.round(rect.width),
                                height: Math.round(rect.height),
                                natural: img.naturalWidth + 'x' + img.naturalHeight,
                                is_visible: rect.width > 0 && rect.height > 0
                            }
                        })
                    """)
                    
                    for img in imgs_desk:
                        src = img['src']
                        if not src or src.startswith('data:'): continue
                        
                        if src not in results: results[src] = {}
                        
                        if img['width'] > 0:
                            results[src]['Desktop'] = f"{img['width']}x{img['height']}"
                            results[src]['Natural'] = img['natural']
                            total_images_found += 1

                except Exception as e:
                    print(f"Desktop Error {url}: {e}")

                try:
                    page_mob.goto(url, wait_until="domcontentloaded", timeout=20000)
                    page_mob.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                    time.sleep(0.5)
                    
                    imgs_mob = page_mob.evaluate("""
                        Array.from(document.querySelectorAll('img')).map(img => {
                            const rect = img.getBoundingClientRect();
                            return { src: img.currentSrc || img.src, width: Math.round(rect.width), height: Math.round(rect.height) }
                        })
                    """)
                    
                    for img in imgs_mob:
                        src = img['src']
                        if not src or src.startswith('data:'): continue
                        if src not in results: results[src] = {}
                        if img['width'] > 0:
                            results[src]['Mobile'] = f"{img['width']}x{img['height']}"

                except Exception:
                    pass

            browser.close()
            return results, total_images_found

    except Exception as e:
        error_details = traceback.format_exc()
        return {}, f"CRASH DETAILS:\n{error_details}"

# --- CRAWLER HANDLERS ---
def crawl_website(start_url, max_urls, crawl_scope, progress_container, ignore_robots=False, custom_selector=None, link_selector=None, search_config=None, use_js=False, storage="RAM"):
    start_url = start_url.strip() # FIX: Ensure the starting URL has no hidden spaces
    crawler = UltraFrogCrawler(max_urls, ignore_robots, crawl_scope, custom_selector, link_selector, search_config, use_js)
    crawler.set_base_url(start_url)
    
    if storage == "SQLite":
        init_db(st.session_state.db_file) 
    else:
        st.session_state.crawl_data = [] 
    
    urls_to_visit = deque([(start_url, 0)]) 
    visited_urls = set()
    st.session_state.total_crawled_count = 0
    
    progress_bar = progress_container.progress(0)
    status_text = progress_container.empty()
    
    worker_count = 1 if use_js else 5
    if use_js: st.warning("🐢 JavaScript Rendering is ON. Speed reduced to 1 URL at a time to prevent crashes.")
    
    with ThreadPoolExecutor(max_workers=worker_count) as executor:
        while urls_to_visit and len(visited_urls) < max_urls:
            if st.session_state.stop_crawling: break
                
            current_batch = []
            batch_size = min(20, len(urls_to_visit), max_urls - len(visited_urls))
            
            for _ in range(batch_size):
                if urls_to_visit and not st.session_state.stop_crawling:
                    url, depth = urls_to_visit.popleft()
                    if url not in visited_urls and crawler.can_fetch(url):
                        current_batch.append((url, depth))
                        visited_urls.add(url)
            
            if not current_batch: break
            
            future_to_url = {executor.submit(crawler.extract_page_data, u): (u, d) for u, d in current_batch}
            
            for future in as_completed(future_to_url):
                if st.session_state.stop_crawling:
                    for f in future_to_url: f.cancel()
                    break
                try:
                    page_data = future.result(timeout=60 if use_js else 12) 
                    _, current_depth = future_to_url[future]
                    page_data['depth'] = current_depth
                    
                    if storage == "SQLite":
                        save_row_to_db(page_data, st.session_state.db_file) 
                    else:
                        st.session_state.crawl_data.append(page_data)

                    st.session_state.total_crawled_count += 1
                    
                    if not st.session_state.stop_crawling:
                        for link_data in page_data.get('internal_links', []):
                            link_url = link_data['url']
                            if (link_url not in visited_urls and 
                                not any(link_url == u[0] for u in urls_to_visit) and 
                                crawler.should_crawl_url(link_url) and 
                                len(visited_urls) + len(urls_to_visit) < max_urls):
                                urls_to_visit.append((link_url, current_depth + 1))
                    
                    progress = min(st.session_state.total_crawled_count / max_urls, 1.0)
                    progress_bar.progress(progress)
                    status_text.text(f"🚀 Crawled: {st.session_state.total_crawled_count} | Queue: {len(urls_to_visit)} | Speed: {st.session_state.total_crawled_count/max(1, time.time() - st.session_state.get('start_time', time.time())):.1f} URLs/sec")
                except Exception as e: st.error(f"Error: {e}")
    return True

def crawl_from_list(url_list, progress_container, ignore_robots=False, custom_selector=None, link_selector=None, search_config=None, use_js=False, storage="RAM"):
    clean_urls = [u.strip() for u in url_list if u.strip()]
    unique_urls = list(set(clean_urls))
    
    if len(clean_urls) > len(unique_urls):
        diff = len(clean_urls) - len(unique_urls)
        st.toast(f"🧹 Removed {diff} duplicate URLs.", icon="ℹ️")

    crawler = UltraFrogCrawler(len(unique_urls), ignore_robots, custom_selector=custom_selector, link_selector=link_selector, search_config=search_config, use_js=use_js)
    
    if storage == "SQLite":
        init_db(st.session_state.db_file) 
    else:
        st.session_state.crawl_data = []

    st.session_state.total_crawled_count = 0
    progress_bar = progress_container.progress(0)
    status_text = progress_container.empty()
    
    valid_urls = [url for url in unique_urls if crawler.can_fetch(url)]
    
    if not valid_urls: 
        st.error("No valid URLs found (check robots.txt or your list).")
        return False
    
    worker_count = 1 if use_js else 10 
    if use_js: st.warning("🐢 JavaScript Rendering is ON. Speed reduced to 1 URL at a time.")

    with ThreadPoolExecutor(max_workers=worker_count) as executor:
        for i in range(0, len(valid_urls), 25):
            if st.session_state.stop_crawling: break
            
            batch = valid_urls[i:i + 25]
            future_to_url = {executor.submit(crawler.extract_page_data, url): url for url in batch}
            
            for future in as_completed(future_to_url):
                url = future_to_url[future]
                
                if st.session_state.stop_crawling:
                    for f in future_to_url: f.cancel()
                    break
                
                try:
                    wait_time = 65 if use_js else 15
                    page_data = future.result(timeout=wait_time)
                    page_data['depth'] = 0 
                    
                    if storage == "SQLite":
                        save_row_to_db(page_data, st.session_state.db_file) 
                    else:
                        st.session_state.crawl_data.append(page_data)

                except TimeoutError:
                    error_data = {
                        'url': url, 'status_code': 0, 'error': 'Timeout - Skipped',
                        'title': 'Skipped', 'depth': 0, 'crawl_timestamp': datetime.now().isoformat()
                    }
                    if storage == "SQLite": save_row_to_db(error_data, st.session_state.db_file)
                    else: st.session_state.crawl_data.append(error_data)

                except Exception as e:
                    error_data = {
                        'url': url, 'status_code': 0, 'error': str(e),
                        'title': 'Error', 'depth': 0, 'crawl_timestamp': datetime.now().isoformat()
                    }
                    if storage == "SQLite": save_row_to_db(error_data, st.session_state.db_file)
                    else: st.session_state.crawl_data.append(error_data)

                st.session_state.total_crawled_count += 1
                progress = st.session_state.total_crawled_count / len(valid_urls)
                progress_bar.progress(min(progress, 1.0))
                status_text.text(f"🚀 Processed: {st.session_state.total_crawled_count}/{len(valid_urls)}")

    return True

def crawl_from_sitemap(sitemap_url, max_urls, progress_container, ignore_robots=False, custom_selector=None, link_selector=None, search_config=None, use_js=False, storage="RAM"):
    crawler = UltraFrogCrawler(max_urls, ignore_robots, custom_selector=custom_selector, link_selector=link_selector, search_config=search_config, use_js=use_js)
    progress_bar = progress_container.progress(0)
    status_text = progress_container.empty()
    status_text.text("🗺️ Extracting URLs from sitemap...")
    sitemap_urls = crawler.extract_sitemap_urls(sitemap_url)
    
    if not sitemap_urls:
        st.error("No URLs found in sitemap")
        return False
    if len(sitemap_urls) > max_urls:
        sitemap_urls = sitemap_urls[:max_urls]
    
    st.info(f"Found {len(sitemap_urls)} URLs in sitemap")
    return crawl_from_list(sitemap_urls, progress_container, ignore_robots, custom_selector, link_selector, search_config, use_js, storage)




# ==========================================
# 1. WORKSPACE SWITCHER (Top of Sidebar)
# ==========================================
with st.sidebar:
    st.markdown("<h2 style='text-align: center; color: #4da8da; font-weight: 800; padding-bottom: 10px;'>🏢 Battersea Agency</h2>", unsafe_allow_html=True)
    workspace = st.selectbox(
        "Select Workspace:",
        ["🔍 SEO & Technical Audit", "📱 Social Media Reporting", "✍️ Content Writers Hub", "🛠️ Other Tools"]
    )
    st.markdown("---")

# ==========================================
# 2. SEO WORKSPACE (Wrap everything here)
# ==========================================
if workspace == "🔍 SEO & Technical Audit":
    with st.sidebar:
        st.markdown("<h2 style='text-align: center; color: #4da8da; font-weight: 800; padding-bottom: 10px;'>⚙️ Settings</h2>", unsafe_allow_html=True)
        
        # 1. Core Config
        with st.container():
            st.markdown("### 🎯 Core Setup")
            storage_option = st.radio(
                "💾 Storage Mode", 
                ["RAM (Fast, <10k URLs)", "SQLite (Unlimited)"], 
                index=0,
                help="Use RAM for speed on small sites. Use SQLite for 100k+ URLs to prevent crashing."
            )

            crawl_mode = st.selectbox("🎯 Crawl Mode", [
                "🕷️ Spider Crawl (Follow Links)",
                "📝 List Mode (Upload URLs)",
                "🗺️ Sitemap Crawl (XML Sitemap)"
            ])
            
            sitemap_url_orphan = ""
            
            if crawl_mode == "🕷️ Spider Crawl (Follow Links)":
                start_url = st.text_input("🌐 Website URL", placeholder="https://example.com")
                sitemap_url_orphan = st.text_input("🗺️ Orphan Check Sitemap (Optional)", placeholder="https://example.com/sitemap.xml")
                max_urls = st.number_input("📊 Max URLs", min_value=1, max_value=100000, value=1000)
                crawl_scope = st.selectbox("🎯 Crawl Scope", ["subfolder", "subdomain", "exact"])
                
            elif crawl_mode == "📝 List Mode (Upload URLs)":
                uploaded_file = st.file_uploader("Choose file", type=['txt', 'csv'])
                url_list_text = st.text_area("Or paste URLs here (one per line)", height=100)
                
            elif crawl_mode == "🗺️ Sitemap Crawl (XML Sitemap)":
                sitemap_url = st.text_input("🗺️ Sitemap URL", placeholder="https://example.com/sitemap.xml")
                max_urls = st.number_input("📊 Max URLs", min_value=1, max_value=100000, value=5000)
                
        # 2. Advanced Config
        with st.expander("🛠️ Advanced Crawler Settings", expanded=False):
            ignore_robots = st.checkbox("🤖 Ignore robots.txt")
            use_js = st.checkbox("🐢 Enable JS Rendering (Playwright)", help="Slow! Use for React/Angular sites only.")
            deep_check = st.checkbox("🐢 Check Everything (Deep Scan)", help="Automatically checks Image Status, File Sizes, and Link Status Codes after crawling.")
            link_selector = st.text_input("Link Area Selector (Optional)", placeholder=".sidebar, #footer, .content", help="Only extract links found inside this CSS element")
            
        # 3. Custom Search Config
        with st.expander("🔍 Custom Search Engine", expanded=False):
            st.caption("Find text, regex, or elements across pages.")
            search_query = st.text_input("Query (Text/Regex/Selector)", help="Enter text to find")
            search_mode = st.selectbox("Search Type", ["Text (Case Insensitive)", "Text (Case Sensitive)", "Regex", "CSS Selector Existence"])
            search_scope = st.selectbox("Look In", ["Visible Text Content", "HTML Source Code"], disabled=(search_mode=="CSS Selector Existence"))
            
        # 4. Custom Extraction Config
        if 'extraction_rules' not in st.session_state:
            st.session_state.extraction_rules = []

        with st.expander("⛏️ Data Extraction Rules", expanded=False):
            st.caption("Add multi-rules like Screaming Frog.")
            
            c_name, c_css = st.columns([1, 2])
            new_name = c_name.text_input("Label", key="new_rule_name", placeholder="Date")
            new_selector = c_css.text_input("CSS Selector", key="new_rule_css", placeholder="meta[property='...']")
            
            c_type, c_attr = st.columns([1, 1])
            new_type = c_type.selectbox("Extract:", ["Text Content", "Attribute Value", "Inner HTML", "HTML Element"], key="new_rule_type")
            new_attr = ""
            if new_type == "Attribute Value":
                new_attr = c_attr.text_input("Attribute", placeholder="content, href", key="new_rule_attr")

            if st.button("➕ Add Rule", use_container_width=True):
                if new_name and new_selector:
                    if new_type == "Attribute Value" and not new_attr:
                        st.error("Attribute Name required.")
                    else:
                        rule = {"name": new_name, "selector": new_selector, "type": new_type, "attribute": new_attr}
                        st.session_state.extraction_rules.append(rule)
                        st.success(f"Added: {new_name}")
                else:
                    st.warning("Label and Selector required.")

            if st.session_state.extraction_rules:
                st.write("---")
                st.write("**Active Rules:**")
                for i, rule in enumerate(st.session_state.extraction_rules):
                    c_desc, c_del = st.columns([4, 1])
                    c_desc.markdown(f"**{rule['name']}**: `{rule['selector']}`")
                    if c_del.button("🗑️", key=f"del_rule_{i}"):
                        st.session_state.extraction_rules.pop(i)
                        st.rerun()
                        
        custom_selector = None # Ensures no crash later 
            
        # 5. Integrations
        with st.expander("🔌 APIs & Integrations", expanded=False):
            st.markdown("#### PageSpeed Insights")
            psi_api_key = st.text_input("Google PSI API Key", type="password")
            
            st.markdown("#### Google Search Console")
            if not HAS_GSC:
                st.error("Missing libraries: google-api-python-client google-auth")
                gsc_property = None
            else:
                LOCAL_KEY_PATH = "gsc_config.json"
                existing_key = None
                if os.path.exists(LOCAL_KEY_PATH):
                    st.success(f"🔑 Key loaded automatically.")
                    with open(LOCAL_KEY_PATH, "r") as f: existing_key = f.read()
                
                gsc_auth_file = st.file_uploader("Upload New JSON Key", type=['json'])
                
                current_file_content = None
                if gsc_auth_file:
                    current_file_content = gsc_auth_file.getvalue().decode("utf-8")
                    with open(LOCAL_KEY_PATH, "w") as f: f.write(current_file_content)
                    st.rerun()
                elif existing_key:
                    current_file_content = existing_key

                if current_file_content:
                    if 'last_gsc_key' not in st.session_state or st.session_state.last_gsc_key != current_file_content:
                        st.session_state.last_gsc_key = current_file_content
                        st.session_state.gsc_service = authenticate_gsc(current_file_content)
                        if st.session_state.gsc_service:
                            st.session_state.gsc_sites_list = list_gsc_sites(st.session_state.gsc_service)
                    
                    if st.session_state.gsc_sites_list:
                        gsc_property = st.selectbox("Select Property", st.session_state.gsc_sites_list)
                        today = datetime.now().date()
                        default_start = today - timedelta(days=28)
                        date_range = st.date_input("📅 Date Range", value=(default_start, today), max_value=today, format="DD/MM/YYYY")
                    else:
                        st.warning("Authenticated, but no sites found.")
                        gsc_property = None
                else:
                    gsc_property = None

        # ACTION BUTTONS
        st.divider()
        col1, col2 = st.columns(2)
        with col1:
            start_btn = st.button("🚀 Start Crawl", type="primary", use_container_width=True, disabled=st.session_state.crawling)
        with col2:
            stop_btn = st.button("⛔ Stop Crawl", use_container_width=True, disabled=not st.session_state.crawling)
        
        st.button("🗑️ Clear All Data", use_container_width=True, on_click=lambda: [
            st.session_state.pop('crawl_data', None),
            st.session_state.pop('psi_results', None),
            os.remove(st.session_state.db_file) if os.path.exists(st.session_state.db_file) else None,
            st.session_state.update({'db_file': f"battersea_data_{uuid.uuid4().hex}.db", 'sitemap_urls_set': set()})
        ])

        if stop_btn:
            st.session_state.stop_crawling = True
            st.session_state.crawling = False
            st.rerun()
        
        if start_btn:
            valid_input = False
            url_list = []
            
            if "SQLite" in storage_option:
                st.session_state.storage_mode = "SQLite"
                st.session_state.crawl_data = [] 
                if os.path.exists(st.session_state.db_file):
                    try: os.remove(st.session_state.db_file)
                    except: pass
                st.session_state.db_file = f"battersea_data_{uuid.uuid4().hex}.db"
            else:
                st.session_state.storage_mode = "RAM"

            search_config = None
            if search_query:
                search_config = {
                    "query": search_query,
                    "mode": search_mode,
                    "scope": search_scope
                }

            if crawl_mode == "🕷️ Spider Crawl (Follow Links)" and start_url:
                valid_input = True
                if sitemap_url_orphan:
                    crawler_temp = UltraFrogCrawler()
                    st.session_state.sitemap_urls_set = set(crawler_temp.extract_sitemap_urls(sitemap_url_orphan))
                else:
                    st.session_state.sitemap_urls_set = set()
            elif crawl_mode == "📝 List Mode (Upload URLs)":
                if uploaded_file:
                    content = uploaded_file.read().decode('utf-8')
                    url_list = [line.strip() for line in content.split('\n') if line.strip()]
                    valid_input = len(url_list) > 0
                elif url_list_text:
                    url_list = [line.strip() for line in url_list_text.split('\n') if line.strip()]
                    valid_input = len(url_list) > 0
            elif crawl_mode == "🗺️ Sitemap Crawl (XML Sitemap)" and sitemap_url:
                valid_input = True
            
            if valid_input:
                st.session_state.crawling = True
                st.session_state.stop_crawling = False
                st.session_state.do_deep_check = deep_check 
                st.session_state.start_time = time.time()
                st.session_state.search_config = search_config
                st.rerun()
            else:
                st.error("Please provide valid input")
            
    def run_post_crawl_deep_check(storage_mode, db_file):
        status_container = st.empty()
        progress_bar = st.progress(0)
        status_container.info("🐢 Deep Check: Gathering URLs and Images...")
        
        all_links = [] 
        all_images = set()
        unique_pages = set()
        
        if storage_mode == "RAM":
            data_source = st.session_state.crawl_data
        else:
            conn = sqlite3.connect(db_file, check_same_thread=False)
            df_temp = pd.read_sql("SELECT url, internal_links, external_links, images FROM pages", conn)
            conn.close()
            data_source = df_temp.to_dict('records')

        for row in data_source:
            unique_pages.add(row['url'])
            for col in ['internal_links', 'external_links']:
                val = row.get(col, [])
                if isinstance(val, str):
                    try: val = json.loads(val)
                    except: val = []
                for link in val:
                    if isinstance(link, dict) and 'url' in link:
                        all_links.append({
                            'Source': row['url'], 
                            'Dest': link['url'], 
                            'Anchor': link.get('anchor_text', '')
                        })
            
            imgs = row.get('images', [])
            if isinstance(imgs, str):
                try: imgs = json.loads(imgs)
                except: imgs = []
            for img in imgs:
                if isinstance(img, dict) and 'src' in img:
                    all_images.add(img['src'])

        links_to_check = list(set([x['Dest'] for x in all_links if x['Dest'] not in st.session_state.link_status_cache]))
        images_to_check = list([u for u in all_images if u not in st.session_state.img_status_cache])
        
        total_ops = len(links_to_check) + len(images_to_check)
        completed = 0
        crawler = UltraFrogCrawler()
        
        if links_to_check:
            status_container.text(f"🐢 Checking Status for {len(links_to_check)} Links...")
            with ThreadPoolExecutor(max_workers=20) as executor:
                future_to_url = {executor.submit(crawler.session.head, u, timeout=5, allow_redirects=True): u for u in links_to_check}
                for future in as_completed(future_to_url):
                    u = future_to_url[future]
                    try:
                        r = future.result()
                        code = r.status_code
                        if code in [405, 403]: 
                            r_get = crawler.session.get(u, timeout=5, stream=True)
                            code = r_get.status_code
                        st.session_state.link_status_cache[u] = code
                    except: st.session_state.link_status_cache[u] = "Error"
                    completed += 1
                    if total_ops > 0: progress_bar.progress(min(completed / total_ops, 1.0))

        import difflib

        if images_to_check:
            status_container.text(f"🖼️ Checking Status & Size for {len(images_to_check)} Images...")
            def check_img(u):
                try:
                    r = crawler.session.head(u, timeout=5, allow_redirects=True)
                    s = r.status_code
                    if s in [405, 403]:
                        r = crawler.session.get(u, timeout=5, stream=True)
                        s = r.status_code
                    size = round(int(r.headers.get('content-length', 0)) / 1024, 2)
                    return u, s, size
                except: return u, "Error", 0

            with ThreadPoolExecutor(max_workers=20) as executor:
                futures = {executor.submit(check_img, u): u for u in images_to_check}
                for f in as_completed(futures):
                    u, code, size = f.result()
                    st.session_state.img_status_cache[u] = code
                    st.session_state.img_size_cache[u] = size
                    completed += 1
                    if total_ops > 0: progress_bar.progress(min(completed / total_ops, 1.0))

        real_dims_to_check = [u for u in all_images if u not in st.session_state.img_real_dim_cache]
        
        if real_dims_to_check:
            status_container.text(f"📏 Measuring Real Dimensions for {len(real_dims_to_check)} Images (Downloading)...")
            prog_pil = 0
            total_pil = len(real_dims_to_check)
            
            def get_pil_dims(u):
                try:
                    r = requests.get(u, timeout=5, headers={'User-Agent': 'Mozilla/5.0'}, verify=False)
                    if r.status_code == 200:
                        img = Image.open(io.BytesIO(r.content))
                        return u, f"{img.width}x{img.height}"
                    return u, f"Error {r.status_code}"
                except: return u, "Error"

            with ThreadPoolExecutor(max_workers=10) as exe:
                futures = {exe.submit(get_pil_dims, u): u for u in real_dims_to_check}
                for f in as_completed(futures):
                    u, dims = f.result()
                    st.session_state.img_real_dim_cache[u] = dims
                    prog_pil += 1
                    progress_bar.progress(min(prog_pil / total_pil, 1.0))

        if HAS_PLAYWRIGHT and unique_pages:
            pages_to_render = [p for p in unique_pages]
            status_container.text(f"👁️ Visual Scan: Rendering {len(pages_to_render)} Pages with Playwright (Slow)...")
            def update_prog_pw(i, total, url):
                progress_bar.progress((i+1)/total)
                status_container.text(f"👁️ Visual Scan: {i+1}/{total} - {url}")

            scan_results, count = measure_rendered_images(pages_to_render, update_prog_pw)
            if isinstance(scan_results, dict):
                st.session_state.img_rendered_cache.update(scan_results)

        status_container.success("✅ Deep Check Complete! All data updated.")
        time.sleep(1)
        status_container.empty()
        
        
    # --- SCHEMA ANALYSIS HELPERS ---
    def validate_schema_structure(schema):
        issues = []
        if not isinstance(schema, dict): return []
        
        s_type = schema.get('@type', '')
        if isinstance(s_type, list): s_type = s_type[0]
        
        if '@context' not in schema: 
            issues.append("⚠️ Missing '@context' (should be 'https://schema.org')")
        
        if s_type == 'Article' or s_type == 'BlogPosting':
            if 'headline' not in schema: issues.append("❌ Missing 'headline'")
            if 'author' not in schema: issues.append("⚠️ Missing 'author'")
            if 'datePublished' not in schema: issues.append("⚠️ Missing 'datePublished'")
            if 'image' not in schema: issues.append("⚠️ Missing 'image' (Required for Google Discover)")

        elif s_type == 'Product':
            if 'name' not in schema: issues.append("❌ Missing 'name'")
            if 'offers' not in schema and 'aggregateRating' not in schema:
                issues.append("❌ Product needs 'offers' (Price) or 'aggregateRating'")
            if 'image' not in schema: issues.append("⚠️ Missing 'image' URL")
        
        elif s_type == 'LocalBusiness' or s_type == 'Organization':
            if 'name' not in schema: issues.append("❌ Missing 'name'")
            if 'address' not in schema: issues.append("⚠️ Missing 'address'")
            if 'telephone' not in schema: issues.append("⚠️ Missing 'telephone'")

        elif s_type == 'BreadcrumbList':
            if 'itemListElement' not in schema: issues.append("❌ Missing 'itemListElement'")

        elif s_type == 'FAQPage':
            if 'mainEntity' not in schema: issues.append("❌ Missing 'mainEntity' (Questions)")

        return issues

    def render_rich_snippet_preview(schema):
        s_type = schema.get('@type', '')
        if isinstance(s_type, list): s_type = s_type[0]
        
        st.markdown("""
        <style>
        .google-card { font-family: arial, sans-serif; background: rgba(128, 128, 128, 0.05); padding: 15px; border-radius: 8px; border: 1px solid rgba(128, 128, 128, 0.2); margin-bottom: 15px; max-width: 600px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); }
        .g-cite { font-size: 12px; line-height: 1.3; opacity: 0.8; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
        .g-title { font-size: 20px; line-height: 1.3; color: #4da8da; margin: 5px 0; text-decoration: none; display: block; cursor: pointer; }
        .g-title:hover { text-decoration: underline; }
        .g-desc { font-size: 14px; line-height: 1.58; opacity: 0.9; }
        .g-meta { font-size: 13px; opacity: 0.7; margin-top: 5px; }
        .g-rating { color: #e7711b; font-weight: bold; }
        </style>
        """, unsafe_allow_html=True)

        title = schema.get('name', schema.get('headline', 'No Title Detected'))
        desc = schema.get('description', 'No description found in schema.')
        
        rating_html = ""
        if 'aggregateRating' in schema:
            try:
                rating = schema['aggregateRating'].get('ratingValue', 4.5)
                count = schema['aggregateRating'].get('reviewCount', 0)
                rating_html = f'<div class="g-meta"><span class="g-rating">★★★★☆ {rating}</span> - {count} reviews</div>'
            except: pass
            
        price_html = ""
        if 'offers' in schema:
            offer = schema['offers']
            if isinstance(offer, list): offer = offer[0]
            price = offer.get('price', '')
            currency = offer.get('priceCurrency', 'USD')
            if price:
                price_html = f'<div class="g-meta" style="font-weight:bold;">{currency} {price}</div>'

        html = f"""
        <div class="google-card">
            <div class="g-cite">https://example.com › ...</div>
            <a class="g-title" href="#">{title}</a>
            <div class="g-desc">{desc[:160]}...</div>
            {rating_html}
            {price_html}
        </div>
        """
        st.markdown(html, unsafe_allow_html=True)
        
    def get_suggested_schema(row):
        suggestions = []
        current = str(row.get('schema_types', '')).lower()
        
        url = str(row.get('url', '')).lower()
        content = str(row.get('scope_content', '')).lower()
        
        if 'product' not in current:
            if '/product/' in url or '/item/' in url: suggestions.append("Product")
            elif 'price' in content and 'add to cart' in content: suggestions.append("Product")
            
        if 'article' not in current and 'blogposting' not in current:
            if '/blog/' in url or '/news/' in url: suggestions.append("Article")
            
        if 'faqpage' not in current:
            if 'frequently asked questions' in content or content.count('?') > 3: suggestions.append("FAQPage")
            
        if 'localbusiness' not in current and 'organization' not in current:
            if 'contact' in url or 'about' in url: suggestions.append("LocalBusiness")
            
        if 'breadcrumblist' not in current:
            suggestions.append("BreadcrumbList")
            
        return ", ".join(suggestions) if suggestions else "✅ Looks Good"


    # --- MAIN CONTENT / DASHBOARD ---
    if st.session_state.crawling:
        st.header("🐸 Battersea Crawler is Running...")
        progress_container = st.container()
        
        try:
            custom_sel = st.session_state.extraction_rules if 'extraction_rules' in st.session_state and st.session_state.extraction_rules else None
            link_sel = link_selector if link_selector else None
            mode_val = st.session_state.storage_mode
            s_conf = st.session_state.get('search_config', None)
            
            if crawl_mode == "🕷️ Spider Crawl (Follow Links)":
                crawl_website(start_url, max_urls, crawl_scope, progress_container, ignore_robots, custom_sel, link_sel, s_conf, use_js, mode_val)
            elif crawl_mode == "📝 List Mode (Upload URLs)":
                if uploaded_file:
                    content = uploaded_file.read().decode('utf-8')
                    url_list = [line.strip() for line in content.split('\n') if line.strip()]
                else:
                    url_list = [line.strip() for line in url_list_text.split('\n') if line.strip()]
                crawl_from_list(url_list, progress_container, ignore_robots, custom_sel, link_sel, s_conf, use_js, mode_val)
            else:
                crawl_from_sitemap(sitemap_url, max_urls, progress_container, ignore_robots, custom_sel, link_sel, s_conf, use_js, mode_val)
            
            if not st.session_state.stop_crawling and st.session_state.get('do_deep_check', False):
                run_post_crawl_deep_check(st.session_state.storage_mode, st.session_state.db_file)

            st.session_state.crawling = False
            st.session_state.stop_crawling = False
            
            if st.session_state.stop_crawling:
                st.warning("⛔ Crawl stopped by user")
            else:
                crawl_time = time.time() - st.session_state.get('start_time', time.time())
                st.success(f"✅ Crawl completed! Found {st.session_state.total_crawled_count} URLs in {crawl_time:.1f} seconds")
            st.rerun()
            
        except Exception as e:
            st.error(f"Error: {str(e)}")
            st.session_state.crawling = False

    # --- RESULT LOADING LOGIC ---
    df = None
    has_data = False

    if st.session_state.storage_mode == "RAM" and st.session_state.crawl_data:
        df = pd.DataFrame(st.session_state.crawl_data)
        has_data = True
    elif st.session_state.storage_mode == "SQLite" and os.path.exists(st.session_state.db_file):
        conn = sqlite3.connect(st.session_state.db_file, check_same_thread=False)
        try:
            df = pd.read_sql("SELECT * FROM pages LIMIT 20000", conn)
            for col in ['internal_links', 'external_links', 'images', 'redirect_chain', 'schema_dump', 'header_structure']: 
                if col in df.columns:
                    df[col] = df[col].apply(lambda x: json.loads(x) if isinstance(x, str) else [])
            has_data = True
            cursor = conn.cursor()
            cursor.execute("SELECT COUNT(*) FROM pages")
            total_in_db = cursor.fetchone()[0]
            if total_in_db > 20000:
                st.warning(f"⚠️ Displaying first 20,000 URLs for performance. Download the DB below for the full dataset ({total_in_db} URLs).")
            conn.close()
        except Exception as e:
            st.error(f"Error reading DB: {e}")

    if has_data and df is not None:
        inlinks_counter = Counter()
        for _, row in df.iterrows():
            links = row.get('internal_links', [])
            if isinstance(links, str): 
                try: links = json.loads(links)
                except: links = []
            for link in links:
                inlinks_counter[link['url']] += 1
                
        df['inlinks_count'] = df['url'].map(inlinks_counter).fillna(0).astype(int)
        
        if 'depth' not in df.columns: df['depth'] = 0
        else: df['depth'] = df['depth'].fillna(0).astype(int)

        if 'custom_search_count' not in df.columns: df['custom_search_count'] = 0
        else: df['custom_search_count'] = df['custom_search_count'].fillna(0).astype(int)

    if has_data and df is not None:
        st.markdown("## 📊 SEO Analysis Dashboard")
        
        col1, col2, col3, col4, col5, col6 = st.columns(6)
        with col1: st.metric("🌐 Total URLs", len(df) if st.session_state.storage_mode == "RAM" else total_in_db)
        with col2: st.metric("✅ Indexable", len(df[df['indexability'] == 'Indexable']))
        with col3: st.metric("❌ Non-Indexable", len(df[df['indexability'] == 'Non-Indexable']))
        with col4: st.metric("🔄 Redirects", len(df[df['redirect_count'] > 0]))
        with col5: st.metric("⚡ Avg Response", f"{df['response_time'].mean():.2f}s" if len(df)>0 else "0s")
        with col6:
            crawled_urls = set(df['url'])
            orphans = list(st.session_state.sitemap_urls_set - crawled_urls) if st.session_state.sitemap_urls_set else []
            st.metric("👻 Orphans", len(orphans))
            
        st.markdown("<br>", unsafe_allow_html=True)
        
        tab1, tab2, tab3, tab_meta_titles, tab_headers, tab_tech, tab10, tab11, tab13, tab14, tab15, tab16, tab_search, tab_cannibal, tab_gsc, tab_audit, tab_competitor = st.tabs([
            "🔗 Internal Links", "🌐 External", "🖼️ Images", "📝 Meta & Titles", "🏗️ Headers", 
            "🛠️ Technical", "📱 Social", "🚀 Performance", 
            "👻 Orphans", "⛏️ Custom Data", "⚡ PSI", "🏗️ Schema", "🔍 Search", "👯 Cannibalization", "📈 GSC", "📅 Audit", "⚔️ Competitor"
        ])
        
        status_lookup = df[['url', 'status_code']].drop_duplicates()
        status_lookup.columns = ['Destination URL', 'Status Code']

        import difflib

        with tab1:
            st.subheader("🔗 Internal Links Analysis")
            
            # --- PART 1: EXISTING LINKS ANALYSIS ---
            if 'internal_links' in df.columns:
                base_df = df[['url', 'internal_links']].copy()
                base_df = base_df.rename(columns={'url': 'Source URL'})
                exploded = base_df.explode('internal_links')
                exploded = exploded.dropna(subset=['internal_links'])
                
                if not exploded.empty:
                    links_data = pd.json_normalize(exploded['internal_links'])
                    exploded = exploded.reset_index(drop=True)
                    links_data = links_data.reset_index(drop=True)
                    final_links = pd.concat([exploded['Source URL'], links_data], axis=1)
                    
                    if 'rel_status' not in final_links.columns: final_links['rel_status'] = 'dofollow'
                    if 'target' not in final_links.columns: final_links['target'] = '_self'
                    
                    final_links = final_links[['Source URL', 'url', 'anchor_text', 'rel_status', 'target', 'placement', 'css_path']]
                    final_links.columns = ['Source URL', 'Destination URL', 'Anchor Text', 'Link Type', 'Target', 'Placement', 'CSS Path']
                    
                    counts_lookup = df[['url', 'inlinks_count', 'internal_links_count']].copy()
                    counts_lookup.columns = ['Source URL', 'Source Inlinks', 'Source Outlinks']
                    final_links = pd.merge(final_links, counts_lookup, on='Source URL', how='left')
                    
                    final_links['Link Scope'] = final_links.apply(
                        lambda x: '🔄 Same Page' if x['Source URL'] == x['Destination URL'] else '➡️ Different Page', 
                        axis=1
                    )

                    final_links = pd.merge(final_links, status_lookup, on='Destination URL', how='left')
                    if 'link_status_cache' not in st.session_state: st.session_state.link_status_cache = {}
                    final_links['Status Code'] = final_links.apply(
                        lambda x: st.session_state.link_status_cache.get(x['Destination URL'], x['Status Code']), axis=1
                    )
                    final_links['Status Code'] = final_links['Status Code'].fillna('Not Crawled').astype(str)

                    # --- TOOLBAR ---
                    st.write("#### 🛠️ Link Tools")
                    c_btn1, c_btn2, c_btn3 = st.columns([1, 1, 1])
                    
                    if c_btn3.button("📊 Calculate Page Authority"):
                        with st.spinner("Calculating PageRank..."):
                            scores = calculate_internal_pagerank(df)
                            st.session_state['pagerank_scores'] = scores
                            st.success("Calculated!")
                            st.rerun()

                    if 'pagerank_scores' in st.session_state:
                        scores = st.session_state['pagerank_scores']
                        final_links['Target Authority'] = final_links['Destination URL'].map(scores).fillna(0).astype(int)

                    if c_btn1.button("🔍 Check Status Codes"):
                        uncrawled = final_links[final_links['Status Code'] == 'Not Crawled']['Destination URL'].unique().tolist()
                        if uncrawled:
                            p_bar = st.progress(0)
                            temp_crawl = UltraFrogCrawler()
                            res = {}
                            def fetch(u):
                                try:
                                    r = temp_crawl.session.head(u, timeout=5)
                                    if r.status_code == 405: r = temp_crawl.session.get(u, timeout=5, stream=True)
                                    return u, r.status_code
                                except: return u, "Error"
                            
                            with ThreadPoolExecutor(max_workers=20) as exc:
                                fut = [exc.submit(fetch, u) for u in uncrawled]
                                for i, f in enumerate(as_completed(fut)):
                                    u, c = f.result()
                                    res[u] = c
                                    if i%5==0: p_bar.progress((i+1)/len(uncrawled))
                            
                            st.session_state.link_status_cache.update(res)
                            st.rerun()

                    col_config = {
                        "Source URL": st.column_config.LinkColumn("From Page"),
                        "Destination URL": st.column_config.LinkColumn("To Page"),
                        "Target Authority": st.column_config.ProgressColumn("Authority", format="%d", min_value=0, max_value=100),
                    }
                    
                    cols_order = ['Source URL', 'Source Inlinks', 'Source Outlinks', 'Destination URL', 'Target Authority', 'Link Scope', 'Anchor Text', 'Status Code', 'CSS Path']
                    existing_cols = [c for c in cols_order if c in final_links.columns]
                    
                    st.dataframe(final_links[existing_cols], use_container_width=True, column_config=col_config)
                    
                    c1, c2, c3, c4 = st.columns(4)
                    c1.metric("Total Links", len(final_links))
                    c2.metric("Self-Referencing", len(final_links[final_links['Link Scope'].str.contains('Same')]))
                    c3.metric("Nofollow", len(final_links[final_links['Link Type'].str.contains('nofollow')]))
                    c4.metric("Broken", len(final_links[final_links['Status Code'].str.match(r'4|5', na=False)]))
                    
                    csv = final_links.to_csv(index=False).encode('utf-8')
                    st.download_button("📥 Download Link Report", csv, "internal_links.csv", "text/csv")

                else:
                    st.warning("No internal links found.")

            # --- PART 2: AI OPPORTUNITIES ---
            st.divider()
            st.write("### 💡 AI Internal Link Opportunities")
            st.info("Find pages that *should* link to each other based on content similarity.")

            if not HAS_SKLEARN:
                st.error("❌ 'scikit-learn' library missing. Install it to use this feature.")
            else:
                c_ai1, c_ai2 = st.columns([1, 1])
                min_score = c_ai1.slider("Minimum Relevance Score", 0, 100, 50)
                max_links = c_ai2.number_input("Max Suggestions per Page", 1, 20, 5)
                
                if st.button("🚀 Generate Link Suggestions", type="primary"):
                    with st.spinner("Analyzing content semantics..."):
                        if 'scope_content' not in df.columns:
                            st.error("Please re-crawl to capture content data.")
                        else:
                            suggestions_df = generate_interlink_suggestions(df, min_score=min_score, max_suggestions=max_links)
                            if not suggestions_df.empty:
                                if 'pagerank_scores' in st.session_state:
                                    scores = st.session_state['pagerank_scores']
                                    suggestions_df['Target Authority'] = suggestions_df['Suggested Target URL'].map(scores).fillna(0).astype(int)
                                
                                st.session_state.interlink_opportunities = suggestions_df
                                st.success(f"Found {len(suggestions_df)} opportunities!")
                            else:
                                st.warning("No suggestions found. Try lowering the score.")

                if 'interlink_opportunities' in st.session_state and not st.session_state.interlink_opportunities.empty:
                    res_df = st.session_state.interlink_opportunities
                    
                    ai_col_config = {
                        "Relevance Score": st.column_config.ProgressColumn("Relevance", format="%.1f%%", min_value=0, max_value=100),
                        "Suggested Target URL": st.column_config.LinkColumn("Target Page"),
                        "Source URL": st.column_config.LinkColumn("Source Page"),
                    }
                    
                    if 'Target Authority' in res_df.columns:
                        ai_col_config["Target Authority"] = st.column_config.ProgressColumn("Target Auth", format="%d", min_value=0, max_value=100)

                    st.dataframe(res_df, column_config=ai_col_config, use_container_width=True)
                    
                    csv_ai = res_df.to_csv(index=False).encode('utf-8')
                    st.download_button("📥 Download Suggestions", csv_ai, "link_opportunities.csv", "text/csv")

        with tab2:
            st.subheader("🌐 External Links Analysis")
            external_data = []
            for _, row in df.iterrows():
                for ext_link in row.get('external_links', []):
                    external_data.append({
                        'Source URL': row['url'],
                        'Destination URL': ext_link['url'],
                        'Anchor Text': ext_link['anchor_text'],
                        'Link Type': ext_link.get('rel_status', 'dofollow'),
                        'Target': ext_link.get('target', '_self'),
                        'Placement': ext_link.get('placement', 'Unknown'),
                        'CSS Path': ext_link.get('css_path', '')
                    })
            if external_data:
                ext_df = pd.DataFrame(external_data)
                ext_df = pd.merge(ext_df, status_lookup, on='Destination URL', how='left')
                
                if 'link_status_cache' not in st.session_state:
                    st.session_state.link_status_cache = {}
                    
                ext_df['Status Code'] = ext_df.apply(
                    lambda x: st.session_state.link_status_cache.get(x['Destination URL'], x['Status Code']), axis=1
                )
                ext_df['Status Code'] = ext_df['Status Code'].fillna('Not Crawled').astype(str)

                col_btn_ext, col_info_ext = st.columns([1, 3])
                uncrawled_ext = ext_df[ext_df['Status Code'] == 'Not Crawled']['Destination URL'].unique().tolist()
                
                if col_btn_ext.button("🔍 Check External Statuses"):
                    if uncrawled_ext:
                        progress_bar = col_info_ext.progress(0)
                        status_text = col_info_ext.empty()
                        temp_crawler = UltraFrogCrawler()
                        results = {}
                        
                        def fetch_status(url):
                            try:
                                r = temp_crawler.session.head(url, timeout=5, allow_redirects=True)
                                if r.status_code == 405 or r.status_code == 403: 
                                    r = temp_crawler.session.get(url, timeout=5, stream=True)
                                return url, r.status_code
                            except Exception as e:
                                return url, "Error"

                        with ThreadPoolExecutor(max_workers=20) as executor:
                            futures = [executor.submit(fetch_status, u) for u in uncrawled_ext]
                            for i, future in enumerate(as_completed(futures)):
                                u, code = future.result()
                                results[u] = code
                                if i % 5 == 0:
                                    progress_bar.progress((i + 1) / len(uncrawled_ext))
                                    status_text.text(f"Checking: {u}")
                        
                        st.session_state.link_status_cache.update(results)
                        status_text.success("✅ External statuses updated!")
                        time.sleep(1)
                        st.rerun()
                    else:
                        col_info_ext.info("All external links have status codes.")

                st.dataframe(ext_df, use_container_width=True)
                csv = ext_df.to_csv(index=False).encode('utf-8')
                st.download_button("📥 Download External", csv, "external_links.csv", "text/csv")
            else:
                st.info("No external links found.")

        with tab3:
            st.subheader("🖼️ Images Analysis")
            
            images_data = []
            for _, row in df.iterrows():
                imgs = row.get('images', [])
                if isinstance(imgs, str):
                    try: imgs = json.loads(imgs)
                    except: imgs = []
                    
                for img in imgs:
                    html_w = str(img.get('width', '')).strip()
                    html_h = str(img.get('height', '')).strip()
                    
                    if html_w and html_h: html_dim_str = f"{html_w}x{html_h}"
                    elif html_w: html_dim_str = f"{html_w}x?"
                    elif html_h: html_dim_str = f"?x{html_h}"
                    else: html_dim_str = "Missing in HTML"

                    images_data.append({
                        'source_url': row['url'],
                        'image_url': img['src'],
                        'alt_text': img['alt'],
                        'html_dimensions': html_dim_str,
                        'real_dimensions': 'Not Checked',
                        'rendered_desktop': 'Not Scanned',
                        'rendered_mobile': 'Not Scanned',
                        'size_kb': img.get('size_kb', 0),
                        'status_code': 'Pending' 
                    })
            
            if images_data:
                img_df = pd.DataFrame(images_data)
                
                if 'img_size_cache' not in st.session_state: st.session_state.img_size_cache = {}
                if st.session_state.img_size_cache:
                    img_df['size_kb'] = img_df['image_url'].map(st.session_state.img_size_cache).fillna(img_df['size_kb'])

                if 'img_real_dim_cache' not in st.session_state: st.session_state.img_real_dim_cache = {}
                if st.session_state.img_real_dim_cache:
                    img_df['real_dimensions'] = img_df.apply(lambda x: st.session_state.img_real_dim_cache.get(x['image_url'], x['real_dimensions']), axis=1)

                if 'img_status_cache' not in st.session_state: st.session_state.img_status_cache = {}
                if st.session_state.img_status_cache:
                    img_df['status_code'] = img_df['image_url'].map(st.session_state.img_status_cache).fillna('Pending')

                if 'img_rendered_cache' not in st.session_state: st.session_state.img_rendered_cache = {}
                
                def normalize_url(u):
                    if not u: return ""
                    u = u.split('?')[0].split('#')[0]
                    return u.replace('https://', '').replace('http://', '').replace('www.', '')

                def get_rendered_data(img_url, key):
                    cache = st.session_state.img_rendered_cache
                    if not cache: return 'Not Scanned'
                    if img_url in cache and key in cache[img_url]: return cache[img_url][key]
                    clean_target = normalize_url(img_url)
                    for cache_url, data in cache.items():
                        if clean_target in cache_url or normalize_url(cache_url) == clean_target:
                            return data.get(key, 'Not Scanned')
                    return 'Not Scanned'

                if st.session_state.img_rendered_cache:
                    img_df['rendered_desktop'] = img_df['image_url'].apply(lambda x: get_rendered_data(x, 'Desktop'))
                    img_df['rendered_mobile'] = img_df['image_url'].apply(lambda x: get_rendered_data(x, 'Mobile'))
                    def update_natural(row):
                        if row['real_dimensions'] == 'Not Checked':
                            return get_rendered_data(row['image_url'], 'Natural').replace('Not Scanned', 'Not Checked')
                        return row['real_dimensions']
                    img_df['real_dimensions'] = img_df.apply(update_natural, axis=1)

                st.write("#### 🛠️ Image Tools")
                col_kb, col_px, col_vis, col_stat = st.columns([1, 1, 1.5, 1])
                
                with col_kb:
                    if st.button("1️⃣ Check File Sizes"):
                        targets = img_df[img_df['size_kb'] == 0]['image_url'].unique().tolist()
                        if targets:
                            p_bar = st.progress(0)
                            crawler_temp = UltraFrogCrawler()
                            res = {}
                            with ThreadPoolExecutor(max_workers=20) as exe:
                                futures = {exe.submit(crawler_temp.get_file_size, u): u for u in targets}
                                for i, f in enumerate(as_completed(futures)):
                                    u = futures[f]
                                    try: res[u] = f.result()
                                    except: res[u] = 0
                                    if i%5==0: p_bar.progress((i+1)/len(targets))
                            st.session_state.img_size_cache.update(res)
                            st.rerun()

                with col_px:
                    if st.button("2️⃣ Check Real Dims"):
                        targets = img_df[
                            (img_df['real_dimensions'] == 'Not Checked') | 
                            (img_df['real_dimensions'] == 'Error')
                        ]['image_url'].unique().tolist()
                        if targets:
                            p_bar = st.progress(0)
                            results_cache = {}
                            def get_pil_dims(u):
                                try:
                                    r = requests.get(u, timeout=8, headers={'User-Agent': 'Mozilla/5.0'}, verify=False)
                                    if r.status_code == 200:
                                        image_file = io.BytesIO(r.content)
                                        img = Image.open(image_file)
                                        return u, f"{img.width}x{img.height}"
                                    return u, f"Error {r.status_code}"
                                except: return u, "Error"
                            
                            with ThreadPoolExecutor(max_workers=10) as exe:
                                futures = {exe.submit(get_pil_dims, u): u for u in targets}
                                for i, future in enumerate(as_completed(futures)):
                                    url = futures[future]
                                    try: results_cache[url] = future.result()[1]
                                    except: results_cache[url] = "Error"
                                    if i % 2 == 0: p_bar.progress((i + 1) / len(targets))
                            st.session_state.img_real_dim_cache.update(results_cache)
                            st.rerun()

                with col_vis:
                    if st.button("3️⃣ Check Visual Dims"):
                        unique_pages = img_df['source_url'].unique().tolist()
                        if not HAS_PLAYWRIGHT: st.error("❌ Playwright not installed.")
                        elif not unique_pages: st.warning("No pages to scan.")
                        else:
                            p_bar = st.progress(0)
                            s_text = st.empty()
                            def update_prog(i, total, url):
                                p_bar.progress((i+1)/total)
                                s_text.text(f"Rendering {i+1}/{total}: {url}")
                            
                            scan_results, img_count = measure_rendered_images(unique_pages, update_prog)
                            if isinstance(scan_results, dict) and scan_results:
                                st.session_state.img_rendered_cache.update(scan_results)
                                st.success(f"✅ Scanned {len(unique_pages)} pages!")
                                time.sleep(1)
                                st.rerun()

                with col_stat:
                    if st.button("4️⃣ Check Status"):
                        targets = img_df[img_df['status_code'].isin(['Pending', 'Error'])]['image_url'].unique().tolist()
                        
                        if targets:
                            p_bar = st.progress(0)
                            stat_text = st.empty()
                            res_status = {}
                            
                            def fetch_img_status(u):
                                try:
                                    r = requests.head(u, timeout=5, headers={'User-Agent': 'Mozilla/5.0'}, verify=False, allow_redirects=True)
                                    if r.status_code in [405, 403]:
                                        r = requests.get(u, timeout=5, headers={'User-Agent': 'Mozilla/5.0'}, stream=True, verify=False)
                                    return u, r.status_code
                                except Exception as e:
                                    return u, "Error"

                            with ThreadPoolExecutor(max_workers=20) as exe:
                                futures = {exe.submit(fetch_img_status, u): u for u in targets}
                                for i, f in enumerate(as_completed(futures)):
                                    u = futures[f]
                                    try: res_status[u] = f.result()[1]
                                    except: res_status[u] = "Error"
                                    
                                    if i % 5 == 0: 
                                        p_bar.progress((i+1)/len(targets))
                                        stat_text.text(f"Checking {i+1}/{len(targets)}")
                            
                            st.session_state.img_status_cache.update(res_status)
                            st.success("✅ Status Checks Complete")
                            time.sleep(0.5)
                            st.rerun()
                        else:
                            st.info("All images checked.")

                st.divider()
                
                def analyze_image_status(row):
                    real = row['real_dimensions']
                    html = row['html_dimensions']
                    vis_d = row['rendered_desktop']
                    code = str(row['status_code'])
                    
                    status = []
                    
                    if code != 'Pending' and code != '200':
                        status.append(f"❌ HTTP {code}")

                    if html == 'Missing in HTML': status.append("⚠️ Missing HTML Attrs")
                    
                    if 'x' in str(real) and 'x' in str(vis_d) and real != 'Not Checked' and vis_d != 'Not Scanned':
                        try:
                            rw, rh = map(int, real.split('x'))
                            vw, vh = map(int, vis_d.split('x'))
                            if vw > 0 and rw < vw: status.append("⚠️ Pixelated (Real < Visible)")
                            if vw > 0 and rw > (vw * 3): status.append("⚠️ Too Big (Real > 3x Visible)")
                        except: pass
                        
                    if not status: return "✅ Good"
                    return " | ".join(status)

                if not img_df.empty:
                    img_df['Analysis'] = img_df.apply(analyze_image_status, axis=1)

                cols_order = ['image_url', 'status_code', 'size_kb', 'source_url', 'alt_text', 'html_dimensions', 'real_dimensions', 'Analysis']
                final_cols = cols_order + [c for c in img_df.columns if c not in cols_order]

                st.dataframe(
                    img_df[final_cols], 
                    use_container_width=True, 
                    column_config={
                        "image_url": st.column_config.LinkColumn("Image Link"),
                        "source_url": st.column_config.LinkColumn("Found On Page"),
                        "size_kb": st.column_config.NumberColumn("KB", format="%.2f"),
                        "status_code": st.column_config.TextColumn("HTTP Status", width="small"),
                    }
                )
                
                csv_img = img_df.to_csv(index=False).encode('utf-8')
                st.download_button("📥 Download Image Report", csv_img, "images_report.csv", "text/csv")

                st.divider()
                st.subheader("📉 Optimize & Resize Images (Bulk & Interactive)")
                
                if not HAS_PIL:
                    st.error("❌ 'Pillow' library missing. Run `pip install Pillow`")
                else:
                    st.info("Identify heavy images from your crawl, shrink them to their actual rendered size, and compress them.")
                    
                    c_f1, c_f2 = st.columns(2)
                    min_kb = c_f1.number_input("Filter Min KB (Find heavy images):", 0, 10000, 100)
                    interactive_seo_mode = c_f2.checkbox("🔍 Interactive Preview Mode (Review 1-by-1)", value=False)
                    
                    # Reset wizard if user unchecks the box
                    if not interactive_seo_mode and 'seo_wizard_active' in st.session_state:
                        st.session_state.seo_wizard_active = False
                        
                    st.divider()
                    
                    candidates = img_df[img_df['size_kb'] >= min_kb].drop_duplicates(subset=['image_url'])
                    target_urls = candidates['image_url'].tolist()

                    if candidates.empty:
                        st.warning(f"No images found larger than {min_kb} KB. Try lowering the filter.")
                        
                    elif not interactive_seo_mode:
                        # ==========================================
                        # STANDARD BULK MODE
                        # ==========================================
                        c1, c2, c3 = st.columns(3)
                        reduc_pct = c1.slider("Quality %", 10, 90, 50, help="Lower = Smaller File")
                        target_fmt_ui = c2.selectbox("Format", ["WebP", "JPEG", "PNG", "Original"])
                        resize_mode = c3.selectbox("Resize To:", ["Original Size Only (No Resize)", "Desktop Rendered Size", "Mobile Rendered Size", "Both Desktop & Mobile"])

                        if st.button("✨ Optimize All (Bulk)", type="primary"):
                            out_dir = "optimized_images"
                            if not os.path.exists(out_dir): os.makedirs(out_dir)

                            progress = st.progress(0)
                            status = st.empty()
                            report_data = []
                            
                            def process_variant(img_obj, url, variant_label, target_w, target_h, original_kb, original_dims_str, old_fmt):
                                try:
                                    if target_w and target_h:
                                        img_obj = img_obj.resize((target_w, target_h), Image.Resampling.LANCZOS)
                                    
                                    new_dims_str = f"{img_obj.width}x{img_obj.height}"
                                    
                                    if target_fmt_ui == "Original": fmt = old_fmt
                                    elif target_fmt_ui == "WebP": fmt = "WEBP"
                                    elif target_fmt_ui == "JPEG": fmt = "JPEG"
                                    elif target_fmt_ui == "PNG": fmt = "PNG"
                                    else: fmt = old_fmt

                                    if img_obj.mode in ("RGBA", "P") and fmt in ["JPEG"]:
                                        img_obj = img_obj.convert("RGB")
                                    
                                    buf = io.BytesIO()
                                    save_args = {'format': fmt, 'optimize': True}
                                    if fmt in ['JPEG', 'WEBP']: save_args['quality'] = reduc_pct
                                    
                                    img_obj.save(buf, **save_args)
                                    new_data = buf.getvalue()
                                    new_size_kb = len(new_data) / 1024

                                    base_name = url.split('/')[-1].split('?')[0].split('.')[0]
                                    if not base_name: base_name = f"img_{uuid.uuid4().hex[:6]}"
                                    base_name = "".join([c for c in base_name if c.isalnum() or c in '_-'])
                                    
                                    suffix = f"_{variant_label.lower()}" if variant_label != "Original" else ""
                                    filename = f"{base_name}{suffix}.{fmt.lower()}"
                                    local_path = os.path.abspath(os.path.join(out_dir, filename))
                                    
                                    with open(local_path, "wb") as f: f.write(new_data)
                                    
                                    return {
                                        "Original URL": url, "Variant": variant_label, "Old Size KB": round(original_kb, 2),
                                        "New Size KB": round(new_size_kb, 2), "Before Dimension": original_dims_str,
                                        "New Dimension": new_dims_str, "Old Format": old_fmt, "New Format": fmt, "Path": local_path
                                    }
                                except Exception:
                                    return None

                            def process_image_row(row_tuple):
                                url = row_tuple.image_url
                                desk_render = row_tuple.rendered_desktop
                                mob_render = row_tuple.rendered_mobile
                                
                                try:
                                    r = requests.get(url, timeout=10, headers={'User-Agent': 'Mozilla/5.0'}, verify=False)
                                    if r.status_code != 200: return []
                                    
                                    original_kb = len(r.content) / 1024
                                    img_org = Image.open(io.BytesIO(r.content))
                                    original_dims_str = f"{img_org.width}x{img_org.height}"
                                    old_fmt = img_org.format if img_org.format else 'JPEG'

                                    results, tasks = [], []
                                    
                                    if resize_mode == "Original Size Only (No Resize)":
                                        tasks.append((None, None, "Original"))
                                        
                                    elif resize_mode == "Desktop Rendered Size":
                                        if desk_render and 'x' in desk_render and desk_render != 'Not Scanned':
                                            dw, dh = map(int, desk_render.split('x'))
                                            tasks.append((dw, dh, "Desktop"))
                                            
                                    elif resize_mode == "Mobile Rendered Size":
                                        if mob_render and 'x' in mob_render and mob_render != 'Not Scanned':
                                            mw, mh = map(int, mob_render.split('x'))
                                            tasks.append((mw, mh, "Mobile"))

                                    elif resize_mode == "Both Desktop & Mobile":
                                        if desk_render and 'x' in desk_render and desk_render != 'Not Scanned':
                                            dw, dh = map(int, desk_render.split('x'))
                                            tasks.append((dw, dh, "Desktop"))
                                        if mob_render and 'x' in mob_render and mob_render != 'Not Scanned':
                                            mw, mh = map(int, mob_render.split('x'))
                                            tasks.append((mw, mh, "Mobile"))

                                    for w, h, label in tasks:
                                        res = process_variant(img_org.copy(), url, label, w, h, original_kb, original_dims_str, old_fmt)
                                        if res: results.append(res)
                                        
                                    return results
                                except Exception: return []

                            with ThreadPoolExecutor(max_workers=5) as exe:
                                futures = [exe.submit(process_image_row, row) for row in candidates.itertuples()]
                                for i, f in enumerate(as_completed(futures)):
                                    res_list = f.result()
                                    if res_list: report_data.extend(res_list)
                                    progress.progress((i + 1) / len(candidates))
                                    status.text(f"Processing: {i+1}/{len(candidates)}")

                            if report_data:
                                st.success(f"✅ Generated {len(report_data)} optimized images in folder: `{out_dir}`")
                                rep_df = pd.DataFrame(report_data)
                                st.dataframe(rep_df, use_container_width=True)
                                csv_rep = rep_df.to_csv(index=False).encode('utf-8')
                                st.download_button("📥 Download Bulk Report", csv_rep, "conversion_report.csv", "text/csv")
                            else:
                                st.warning("No images processed successfully.")

                    else:
                        # ==========================================
                        # NEW INTERACTIVE SEO WIZARD
                        # ==========================================
                        if 'seo_wizard_active' not in st.session_state:
                            st.session_state.seo_wizard_active = False

                        if not st.session_state.seo_wizard_active:
                            st.write(f"**Found {len(target_urls)} images.** Ready to review them one by one?")
                            if st.button("🚀 Start Interactive Wizard", type="primary", use_container_width=True):
                                st.session_state.seo_wizard_active = True
                                st.session_state.seo_int_img_index = 0
                                st.session_state.seo_int_processed_imgs = {}
                                st.session_state.seo_last_target_urls = target_urls
                                st.rerun()
                        else:
                            # State management specifically for the SEO tab
                            if 'seo_int_img_index' not in st.session_state:
                                st.session_state.seo_int_img_index = 0
                            if 'seo_int_processed_imgs' not in st.session_state:
                                st.session_state.seo_int_processed_imgs = {}

                            # Reset if the list of filtered URLs changes
                            if target_urls != st.session_state.seo_last_target_urls:
                                st.session_state.seo_last_target_urls = target_urls
                                st.session_state.seo_int_img_index = 0
                                st.session_state.seo_int_processed_imgs = {}

                            if st.session_state.seo_int_img_index < len(target_urls):
                                current_idx = st.session_state.seo_int_img_index
                                current_url = target_urls[current_idx]
                                row_data = candidates[candidates['image_url'] == current_url].iloc[0]
                                
                                st.write(f"### Image {current_idx + 1} of {len(target_urls)}")
                                st.markdown(f"**Found on:** `{row_data['source_url']}`")
                                st.markdown(f"**Image Link:** `{current_url}`")
                                
                                desk_render = row_data.get('rendered_desktop', 'Not Scanned')
                                mob_render = row_data.get('rendered_mobile', 'Not Scanned')
                                st.info(f"💡 **Playwright Detected Sizes** -> Desktop: `{desk_render}` | Mobile: `{mob_render}`")

                                try:
                                    # 1. Download the Live Image
                                    r = requests.get(current_url, timeout=10, headers={'User-Agent': 'Mozilla/5.0'}, verify=False)
                                    if r.status_code == 200:
                                        orig_bytes = r.content
                                        orig_img = Image.open(io.BytesIO(orig_bytes))
                                        orig_kb = len(orig_bytes) / 1024
                                        
                                        # 2. Interactive Controls
                                        c_ctrl1, c_ctrl2, c_ctrl3 = st.columns(3)
                                        dyn_width = c_ctrl1.number_input("Max Width (px)", value=min(1200, orig_img.width), step=50, key=f"seo_w_{current_idx}")
                                        dyn_quality = c_ctrl2.slider("Quality %", 10, 100, 75, key=f"seo_q_{current_idx}")
                                        dyn_fmt = c_ctrl3.selectbox("Format", ["WEBP", "JPEG", "PNG"], key=f"seo_fmt_{current_idx}")
                                        
                                        # 3. Process Live
                                        if orig_img.mode in ("RGBA", "P") and dyn_fmt in ["JPEG"]: 
                                            proc_img = orig_img.convert("RGB")
                                        else:
                                            proc_img = orig_img.copy()
                                            
                                        if proc_img.width > dyn_width:
                                            ratio = dyn_width / float(proc_img.width)
                                            new_h = int((float(proc_img.height) * float(ratio)))
                                            proc_img = proc_img.resize((dyn_width, new_h), Image.Resampling.LANCZOS)
                                            
                                        img_byte_arr = io.BytesIO()
                                        save_args = {'format': dyn_fmt}
                                        if dyn_fmt in ['JPEG', 'WEBP']: save_args['quality'] = dyn_quality
                                        if dyn_fmt in ['JPEG', 'WEBP', 'PNG']: save_args['optimize'] = True
                                            
                                        proc_img.save(img_byte_arr, **save_args)
                                        new_bytes = img_byte_arr.getvalue()
                                        new_kb = len(new_bytes) / 1024
                                        
                                        # 4. Preview
                                        col_l, col_r = st.columns(2)
                                        with col_l:
                                            st.markdown(f"**Original:** {orig_img.width}x{orig_img.height}px | **{orig_kb:.1f} KB**")
                                            st.image(orig_img, use_container_width=True)
                                        with col_r:
                                            savings = 100 - ((new_kb / orig_kb) * 100) if orig_kb > 0 else 0
                                            color = "#38ef7d" if savings > 0 else "#ff4b4b"
                                            st.markdown(f"**Preview:** {proc_img.width}x{proc_img.height}px | **{new_kb:.1f} KB** (<span style='color:{color}; font-weight:bold;'>-{savings:.1f}%</span>)", unsafe_allow_html=True)
                                            st.image(new_bytes, use_container_width=True)
                                        
                                        st.write("")
                                        
                                        # 5. Buttons
                                        c_btn1, c_btn2, c_btn3 = st.columns([1, 1, 2])
                                        if c_btn1.button("✅ Accept & Next", type="primary", use_container_width=True):
                                            base_name = current_url.split('/')[-1].split('?')[0].split('.')[0]
                                            if not base_name: base_name = f"img_{uuid.uuid4().hex[:6]}"
                                            base_name = "".join([c for c in base_name if c.isalnum() or c in '_-'])
                                            save_name = f"{base_name}_opt.{dyn_fmt.lower()}"
                                            
                                            st.session_state.seo_int_processed_imgs[save_name] = new_bytes
                                            st.session_state.seo_int_img_index += 1
                                            st.rerun()
                                            
                                        if c_btn2.button("⏭️ Skip Image", use_container_width=True):
                                            st.session_state.seo_int_img_index += 1
                                            st.rerun()
                                            
                                    else:
                                        st.error(f"Failed to fetch image from URL (HTTP {r.status_code}). Target site might block bots.")
                                        if st.button("⏭️ Skip Blocked Image"):
                                            st.session_state.seo_int_img_index += 1
                                            st.rerun()
                                            
                                except Exception as e:
                                    st.error(f"Error loading image: {e}")
                                    if st.button("⏭️ Skip Broken Image"):
                                        st.session_state.seo_int_img_index += 1
                                        st.rerun()
                                        
                            else:
                                st.success(f"🎉 You have reviewed all {len(target_urls)} flagged images!")
                                
                                if st.session_state.seo_int_processed_imgs:
                                    st.write(f"You optimized **{len(st.session_state.seo_int_processed_imgs)}** images.")
                                    zip_buf = io.BytesIO()
                                    with zipfile.ZipFile(zip_buf, "a", zipfile.ZIP_DEFLATED, False) as zf:
                                        for name, _bytes in st.session_state.seo_int_processed_imgs.items():
                                            zf.writestr(name, _bytes)
                                    
                                    st.download_button("📥 Download Final ZIP", zip_buf.getvalue(), "seo_interactive_optimized.zip", "application/zip", type="primary")
                                else:
                                    st.warning("You skipped all images. Nothing to download.")
                                
                                if st.button("🔄 Close Wizard"):
                                    st.session_state.seo_wizard_active = False
                                    st.rerun()

        with tab_meta_titles:
            st.subheader("📝 Meta Tags & Titles Analysis")
            meta_view = df[['url', 'title', 'title_length', 'meta_description', 'meta_desc_length', 'h1_tags']].copy()
            st.dataframe(meta_view, use_container_width=True)
            csv_meta = meta_view.to_csv(index=False).encode('utf-8')
            st.download_button("📥 Download All Meta Data", csv_meta, "meta_titles.csv", "text/csv")

            st.divider()
            with st.expander("✨ AI Content Generator (Titles & Meta)", expanded=False):
                st.info("Generate missing or improved tags using AI.")
                
                c1, c2 = st.columns(2)
                provider = c1.selectbox("Provider", ["Gemini", "OpenAI Compatible (Groq/Ollama/OpenAI)"])
                api_key_gen = c2.text_input("API Key (Leave empty for Ollama)", type="password")
                
                c3, c4 = st.columns(2)
                model_name = c3.text_input("Model Name", value="gemini-1.5-flash" if provider=="Gemini" else "lama-3.1-8b-instant")
                endpoint = c4.text_input("Endpoint URL (OpenAI/Ollama Only)", value="https://api.groq.com/openai/v1/chat/completions")
                
                action_type = st.radio("I want to generate:", ["Titles Only", "Meta Descriptions Only"], horizontal=True)
                filter_mode = st.radio("Generate for:", ["Only Missing Items (Empty)", "Regenerate All"], horizontal=True)
                
                if st.button("🚀 Generate Content", type="primary"):
                    if action_type == "Titles Only":
                        if filter_mode == "Only Missing Items (Empty)":
                            targets = meta_view[meta_view['title_length'] == 0].copy()
                        else:
                            targets = meta_view.copy()
                    else: 
                        if filter_mode == "Only Missing Items (Empty)":
                            targets = meta_view[meta_view['meta_desc_length'] == 0].copy()
                        else:
                            targets = meta_view.copy()
                    
                    if targets.empty:
                        st.warning("No pages match your criteria.")
                    else:
                        st.write(f"Generating for {len(targets)} pages...")
                        progress_gen = st.progress(0)
                        results_gen = []
                        
                        def process_gen_row(row):
                            content_snippet = str(row.get('scope_content', df[df['url'] == row['url']]['scope_content'].values[0]))[:2000] 
                            
                            if action_type == "Titles Only":
                                context = f"Current H1: {row.get('h1_tags', '')}\nPage Content: {content_snippet}"
                                prompt = f"""
                                Act as a Senior SEO Copywriter. Generate ONE highly optimized SEO Title tag for this webpage.
                                
                                STRICT RULES:
                                1. Length: Must be exactly between 50 and 60 characters.
                                2. Intent: Make it highly clickable (High CTR) but DO NOT use cheap clickbait.
                                3. Keywords: Put the main topic/keyword near the front.
                                4. Format: Output ONLY the title text. Do NOT wrap it in quotes. Do NOT say "Here is your title".
                                
                                Context:
                                {context}
                                """
                                col_name = "New AI Title"
                                old_val = row['title']
                            else:
                                context = f"Current Title: {row.get('title', '')}\nPage Content: {content_snippet}"
                                prompt = f"""
                                Act as a Senior SEO Copywriter. Generate ONE highly optimized SEO Meta Description for this webpage.
                                
                                STRICT RULES:
                                1. Length: Must be exactly between 145 and 155 characters.
                                2. Intent: Clearly explain the exact value the user will get by clicking.
                                3. CTA: End with a natural, subtle Call-to-Action (e.g., Learn more, Discover how, Get started).
                                4. Tone: Professional, engaging, and active voice.
                                5. Format: Output ONLY the description text. Do NOT wrap it in quotes. Do NOT add any extra conversational text.
                                
                                Context:
                                {context}
                                """
                                col_name = "New AI Description"
                                old_val = row['meta_description']
                            
                            generated = generate_ai_meta(provider, api_key_gen, model_name, endpoint, prompt, "You are a strict Technical SEO Expert. Follow length rules perfectly.")
                            
                            clean_generated = generated.strip().strip('"').strip("'").replace("**", "")
                            
                            return {"URL": row['url'], "Old Value": old_val, col_name: clean_generated}

                        with ThreadPoolExecutor(max_workers=5) as executor:
                            futures = [executor.submit(process_gen_row, row) for _, row in targets.iterrows()]
                            for i, f in enumerate(as_completed(futures)):
                                res = f.result()
                                results_gen.append(res)
                                progress_gen.progress((i + 1) / len(targets))
                        
                        res_df = pd.DataFrame(results_gen)
                        st.success("✅ Generation Complete!")
                        st.dataframe(res_df, use_container_width=True)
                        csv_gen = res_df.to_csv(index=False).encode('utf-8')
                        st.download_button("📥 Download Generated Data", csv_gen, "ai_generated_tags.csv", "text/csv")

        with tab_headers:
            st.subheader("🏗️ Header Architecture & Counts")
            
            if 'header_structure' in df.columns:
                struct_df = df[['url', 'h1_count', 'h2_count', 'h3_count', 'h4_count', 'h1_tags', 'header_structure']].copy()
                
                analyzed_data = []
                problematic_urls = []
                bad_h1_count = 0
                broken_struct_count = 0
                
                def get_struct_status(row):
                    struct = row['header_structure']
                    if isinstance(struct, str):
                        try: struct = json.loads(struct)
                        except: struct = []
                    
                    issues, h1_c = analyze_heading_structure(struct)
                    
                    has_h1_issue = h1_c != 1
                    has_hierarchy_issue = any("Skipped" in i for i in issues)
                    
                    status_label = "✅ Perfect"
                    if has_h1_issue and has_hierarchy_issue: status_label = "❌ H1 & Hierarchy Errors"
                    elif has_h1_issue: status_label = "❌ Bad H1 Count"
                    elif has_hierarchy_issue: status_label = "⚠️ Skipped Levels"
                    
                    analyzed_data.append({
                        'url': row['url'],
                        'structure': struct,
                        'issues': issues,
                        'h1_count': h1_c,
                        'status': status_label
                    })
                    
                    if has_h1_issue or has_hierarchy_issue:
                        problematic_urls.append({
                            'URL': row['url'],
                            'H1 Count': h1_c,
                            'Hierarchy Issues': " | ".join(issues)
                        })
                        
                    return status_label

                struct_df['Hierarchy Status'] = struct_df.apply(get_struct_status, axis=1)
                
                bad_h1_count = len(struct_df[struct_df['h1_count'] != 1])
                broken_struct_count = len(struct_df[struct_df['Hierarchy Status'].str.contains("Skipped")])

                m1, m2, m3, m4 = st.columns(4)
                m1.metric("Total Pages", len(df))
                m2.metric("❌ Bad H1 Usage", bad_h1_count, help="Pages with 0 or >1 H1 tags")
                m3.metric("⚠️ Broken Levels", broken_struct_count, help="Pages skipping levels (e.g. H2->H4)")
                m4.metric("✅ Perfect Structure", len(df) - len(problematic_urls))

                st.divider()

                st.write("### 📊 Overview Table")
                cols_to_show = ['Hierarchy Status', 'url', 'h1_count', 'h2_count', 'h3_count', 'h1_tags']
                st.dataframe(
                    struct_df[cols_to_show],
                    use_container_width=True,
                    column_config={
                        "url": st.column_config.LinkColumn("Page URL"),
                        "h1_tags": st.column_config.TextColumn("H1 Text", width="large"),
                        "Hierarchy Status": st.column_config.TextColumn("Status", width="medium"),
                    }
                )
                
                csv = struct_df.to_csv(index=False).encode('utf-8')
                st.download_button("📥 Download Header Report", csv, "header_analysis.csv", "text/csv")

                st.divider()

                c_insp, c_tree = st.columns([1, 2])
                
                with c_insp:
                    st.write("### 🔍 Visual Inspector")
                    st.info("Select a page to visualize its Heading Tree (DOM).")
                    
                    filter_mode = st.radio("Filter List:", ["All Pages", "Only Pages with Issues"], horizontal=True)
                    
                    if filter_mode == "Only Pages with Issues":
                        options = [d['URL'] for d in problematic_urls]
                    else:
                        options = [d['url'] for d in analyzed_data]
                    
                    if not options:
                        st.warning("No pages match this filter.")
                        selected_page = None
                    else:
                        selected_page = st.selectbox("Select Page:", options, key="merged_struct_select")

                with c_tree:
                    if selected_page:
                        page_data = next((item for item in analyzed_data if item['url'] == selected_page), None)
                        
                        if page_data:
                            st.markdown(f"**Analysis for:** `{selected_page}`")
                            
                            if page_data['issues']:
                                for issue in page_data['issues']:
                                    if "❌" in issue: st.error(issue)
                                    else: st.warning(issue)
                            else:
                                st.success("✅ Structure is perfectly logical.")

                            st.markdown("#### 🌳 Heading Tree")
                            st.markdown("""
                            <style>
                            .header-node { padding: 3px 8px; border-left: 3px solid rgba(128,128,128,0.3); margin-bottom: 3px; font-family: monospace; }
                            .h1-node { border-left-color: #ff4b4b; background-color: rgba(255, 75, 75, 0.1); font-weight: bold; }
                            .h2-node { border-left-color: #ffbd45; background-color: rgba(255, 189, 69, 0.1); }
                            .h3-node { border-left-color: #92c5de; }
                            .h4-node { border-left-color: rgba(128,128,128,0.5); opacity: 0.8; }
                            </style>
                            """, unsafe_allow_html=True)

                            if not page_data['structure']:
                                st.info("No headers found on this page.")
                            
                            for h in page_data['structure']:
                                lvl = h['level']
                                indent = (lvl - 1) * 20
                                css_class = "header-node"
                                if lvl == 1: css_class += " h1-node"
                                elif lvl == 2: css_class += " h2-node"
                                elif lvl == 3: css_class += " h3-node"
                                elif lvl >= 4: css_class += " h4-node"
                                
                                st.markdown(
                                    f"<div class='{css_class}' style='margin-left: {indent}px;'>"
                                    f"<span>{h['tag'].upper()}</span>: {h['text']}"
                                    f"</div>", 
                                    unsafe_allow_html=True
                                )
            else:
                st.warning("Header structure data not available. Please re-crawl the site.")

        with tab_tech:
            st.header("🛠️ Master Technical Audit")
            st.info("Unified view of Status Codes, Redirects, and Canonicals.")

            tech_df = df[['url', 'status_code', 'indexability', 'redirect_count', 'final_url', 'canonical_url', 'meta_robots']].copy()

            def get_canonical_status(row):
                c_url = row['canonical_url']
                if not c_url: return "❌ Missing"
                if row['url'] == c_url: return "✅ Self-Referencing"
                return "⚠️ Canonicalized to Other"

            tech_df['Canonical Status'] = tech_df.apply(get_canonical_status, axis=1)

            def get_redirect_status(row):
                if row['redirect_count'] > 0:
                    return f"🔄 Redirect ({row['redirect_count']} hops)"
                if row['status_code'] == 200:
                    return "✅ 200 OK"
                return f"⚠️ {row['status_code']}"

            tech_df['Page Status'] = tech_df.apply(get_redirect_status, axis=1)

            display_df = tech_df[[
                'url', 
                'Page Status', 
                'indexability', 
                'final_url', 
                'Canonical Status', 
                'canonical_url',
                'meta_robots'
            ]]

            display_df.columns = [
                'Source URL', 'Status Overview', 'Indexability', 
                'Redirect Target (Final URL)', 'Canonical Status', 'Canonical Link', 'Robots Tag'
            ]

            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Total URLs", len(tech_df))
            m2.metric("🔄 Redirects", len(tech_df[tech_df['redirect_count'] > 0]))
            m3.metric("❌ Broken (4xx/5xx)", len(tech_df[tech_df['status_code'] >= 400]))
            m4.metric("⚠️ Canonical Issues", len(tech_df[tech_df['Canonical Status'].str.contains("Missing|Other")]))

            st.divider()

            st.dataframe(
                display_df,
                use_container_width=True,
                column_config={
                    "Source URL": st.column_config.LinkColumn("Source URL", width="medium"),
                    "Redirect Target (Final URL)": st.column_config.LinkColumn("Redirect Target"),
                    "Canonical Link": st.column_config.LinkColumn("Canonical Tag"),
                    "Status Overview": st.column_config.TextColumn("Status", width="small"),
                }
            )

            csv = display_df.to_csv(index=False).encode('utf-8')
            st.download_button(
                "📥 Download Full Technical Report", 
                csv, 
                "master_technical_audit.csv", 
                "text/csv",
                type="primary"
            )

        with tab10:
            st.subheader("📱 Social Tags")
            st.dataframe(df[['url', 'og_title', 'twitter_title']], use_container_width=True)
            csv = df[['url', 'og_title', 'twitter_title']].to_csv(index=False).encode('utf-8')
            st.download_button("📥 Download Social", csv, "social_tags.csv", "text/csv")

        with tab11:
            st.subheader("🚀 Performance Stats")
            st.dataframe(df[['url', 'response_time', 'word_count', 'content_length']], use_container_width=True)
            csv = df[['url', 'response_time']].to_csv(index=False).encode('utf-8')
            st.download_button("📥 Download Performance", csv, "performance.csv", "text/csv")

        with tab13:
            st.subheader("👻 Orphan Pages")
            if orphans:
                orphan_df = pd.DataFrame(orphans, columns=['Orphan URL'])
                st.dataframe(orphan_df, use_container_width=True)
                csv = orphan_df.to_csv(index=False).encode('utf-8')
                st.download_button("📥 Download Orphans", csv, "orphans.csv", "text/csv")
            else:
                st.success("No orphans found (or no sitemap provided).")

        with tab14:
            st.subheader("⛏️ Custom Extracted Data")
            
            if 'custom_extraction' not in df.columns:
                st.info("No custom data found. Add rules in the sidebar and start a crawl.")
            else:
                try:
                    valid_custom = df[df['custom_extraction'].notna() & (df['custom_extraction'] != '')].copy()
                    
                    if valid_custom.empty:
                        st.warning("No data extracted yet.")
                    else:
                        json_data = valid_custom['custom_extraction'].apply(
                            lambda x: json.loads(x) if isinstance(x, str) and x.startswith('{') else {}
                        ).tolist()
                        
                        extracted_df = pd.DataFrame(json_data)
                        result_df = pd.concat([valid_custom['url'].reset_index(drop=True), extracted_df], axis=1)
                        
                        st.write(f"**Found data for {len(result_df)} URLs**")
                        st.dataframe(result_df, use_container_width=True)
                        
                        csv = result_df.to_csv(index=False).encode('utf-8')
                        st.download_button("📥 Download Custom Data", csv, "custom_extraction.csv", "text/csv")
                except Exception as e:
                    st.error(f"Error parsing data: {e}")
                    st.dataframe(df[['url', 'custom_extraction']], use_container_width=True)

        with tab15:
            st.subheader("⚡ Google PageSpeed Insights (PSI)")
            st.info("Enter your Google PageSpeed API Key in the Sidebar to use this feature.")
            
            if psi_api_key:
                check_all = st.checkbox("Select All URLs (Run for everyone)", help="⚠️ Be careful with API limits if you have thousands of URLs.")
                if check_all:
                    urls_to_test = df['url'].tolist()
                    st.write(f"Selected {len(urls_to_test)} URLs.")
                else:
                    urls_to_test = st.multiselect("Select URLs to Test (Max 5 recommended)", df['url'].head(20).tolist())
                
                if st.button("🏃 Run PageSpeed Test (Mobile & Desktop)", type="primary"):
                    if not urls_to_test:
                        st.warning("Please select at least one URL.")
                    else:
                        progress_psi = st.progress(0)
                        status_text = st.empty()
                        results = []
                        total_steps = len(urls_to_test)
                        
                        for i, u in enumerate(urls_to_test):
                            status_text.text(f"Testing {i+1}/{total_steps}: {u}...")
                            mobile_res = run_psi_test(u, psi_api_key, "mobile")
                            desktop_res = run_psi_test(u, psi_api_key, "desktop")
                            
                            row = {'url': u}
                            if "error" in mobile_res: row['Mobile Error'] = mobile_res['error']
                            else:
                                for k, v in mobile_res.items(): row[f"Mobile {k}"] = v
                                    
                            if "error" in desktop_res: row['Desktop Error'] = desktop_res['error']
                            else:
                                for k, v in desktop_res.items(): row[f"Desktop {k}"] = v

                            results.append(row)
                            progress_psi.progress((i + 1) / total_steps)
                        
                        st.session_state.psi_results = results
                        status_text.success("✅ Analysis Complete!")
                
                if st.session_state.psi_results:
                    psi_df = pd.DataFrame(st.session_state.psi_results)
                    desired_order = ['url', 'Mobile Score', 'Desktop Score', 'Mobile LCP', 'Desktop LCP', 'Mobile FCP', 'Desktop FCP', 'Mobile CLS', 'Desktop CLS', 'Mobile INP', 'Desktop INP']
                    final_cols = [c for c in desired_order if c in psi_df.columns]
                    remaining_cols = [c for c in psi_df.columns if c not in final_cols]
                    st.dataframe(psi_df[final_cols + remaining_cols], use_container_width=True)
                    csv_psi = psi_df.to_csv(index=False).encode('utf-8')
                    st.download_button("📥 Download PSI Report", csv_psi, "psi_report.csv", "text/csv")
            else:
                st.warning("⚠️ PSI API Key is missing. Please add it in the sidebar.")

        with tab16:
            st.subheader("🏗️ Schema Markup Analysis")

            schema_view = df[['url', 'schema_types', 'schema_validity']].copy()
            schema_view['Rule-Based Suggestion'] = df.apply(get_suggested_schema, axis=1)

            with st.expander("✨ Get AI-Powered Schema Suggestions"):
                st.info("Use AI to read the actual page content and recommend the exact Schema types you should add.")
                c_ai1, c_ai2 = st.columns(2)
                sch_provider = c_ai1.selectbox("AI Provider", ["Gemini", "OpenAI Compatible (Groq/Ollama/OpenAI)"], key="sch_p")
                sch_key = c_ai2.text_input("API Key (or empty for Ollama)", type="password", key="sch_k")
                
                c_ai3, c_ai4 = st.columns(2)
                sch_model = c_ai3.text_input("Model", value="gemini-1.5-flash" if sch_provider=="Gemini" else "llama-3.3-70b-versatile", key="sch_m")
                sch_url = c_ai4.text_input("Endpoint URL", value="https://api.groq.com/openai/v1/chat/completions", key="sch_u")

                if st.button("🧠 Check Schema Suggestion with AI"):
                    with st.spinner("AI is analyzing content to suggest schemas..."):
                        ai_results = {}
                        progress_sch = st.progress(0)
                        total_urls = len(df)
                        
                        def fetch_ai_schema(row):
                            content = str(row.get('scope_content', ''))[:1500] 
                            title = str(row.get('title', ''))
                            
                            prompt = f"""
                            Analyze this webpage content.
                            Title: {title}
                            Content Snippet: {content}
                            
                            Task: Suggest the top 1-3 most relevant Schema.org markup types for this page (e.g., Article, FAQPage, Product, LocalBusiness, Recipe, JobPosting, SoftwareApplication). 
                            ONLY output the exact schema names separated by a comma. Do not include any other text, explanations, or code.
                            """
                            
                            res = generate_ai_meta(sch_provider, sch_key, sch_model, sch_url, prompt, "You are a Technical SEO Expert.")
                            clean_res = res.replace("DECISION:", "").replace("`", "").strip()
                            return row['url'], clean_res

                        with ThreadPoolExecutor(max_workers=3) as executor:
                            futures = {executor.submit(fetch_ai_schema, row): row for _, row in df.iterrows()}
                            for i, f in enumerate(as_completed(futures)):
                                url, suggestion = f.result()
                                ai_results[url] = suggestion
                                progress_sch.progress((i + 1) / total_urls)
                        
                        st.session_state.ai_schema_suggestions = ai_results
                        st.success("✅ AI Suggestions Generated!")

            if 'ai_schema_suggestions' in st.session_state:
                schema_view['✨ AI Suggestion'] = schema_view['url'].map(st.session_state.ai_schema_suggestions)
            else:
                schema_view['✨ AI Suggestion'] = "Click AI Button ⬆️"

            st.write("### 📊 Schema Overview")
            st.info("Select a row to see the Google Preview and Validation details below.")

            col_config = {
                "url": st.column_config.LinkColumn("Page URL", width="medium"),
                "schema_types": st.column_config.TextColumn("Detected Schema", width="small"),
                "Rule-Based Suggestion": st.column_config.TextColumn("Standard Suggestion", width="small"),
                "schema_validity": st.column_config.TextColumn("Status", width="small"),
                "✨ AI Suggestion": st.column_config.TextColumn("✨ AI Suggestion", width="medium")
            }

            selection = st.dataframe(
                schema_view,
                use_container_width=True,
                column_config=col_config,
                on_select="rerun", 
                selection_mode="single-row"
            )

            st.divider()
            
            if selection.selection.rows:
                idx = selection.selection.rows[0]
                selected_url = schema_view.iloc[idx]['url']
                
                st.markdown(f"### 🔍 Inspecting: `{selected_url}`")
                
                row = df[df['url'] == selected_url].iloc[0]
                schema_dump = row.get('schema_dump', [])
                
                col_preview, col_json = st.columns([1, 1])
                
                with col_preview:
                    st.markdown("#### 📱 Google Search Preview")
                    
                    if isinstance(schema_dump, str):
                        try: schema_objs = json.loads(schema_dump)
                        except: schema_objs = []
                    else: schema_objs = schema_dump
                    
                    if not schema_objs:
                        st.warning("⚠️ No Schema JSON found to render preview.")
                        st.markdown("""
                        <div style="font-family:arial; background:rgba(128,128,128,0.05); padding:15px; border:1px solid rgba(128,128,128,0.2); border-radius:8px;">
                            <div style="font-size:12px; opacity:0.8;">https://example.com › ...</div>
                            <div style="font-size:20px; color:#4da8da; margin:5px 0;">Page Title Example</div>
                            <div style="font-size:14px; opacity:0.9;">No rich snippets detected. This is how a standard result looks.</div>
                        </div>
                        """, unsafe_allow_html=True)
                    else:
                        render_rich_snippet_preview(schema_objs[0] if isinstance(schema_objs, list) else schema_objs)
                        
                        st.markdown("#### ✅ Validation Issues")
                        issues = validate_schema_structure(schema_objs[0] if isinstance(schema_objs, list) else schema_objs)
                        if issues:
                            for i in issues: st.error(i)
                        else:
                            st.success("Structure looks valid for Google Rich Results.")

                    encoded_url = requests.utils.quote(selected_url)
                    st.markdown(f"""
                    <div style="display: flex; gap: 10px; margin-top: 20px;">
                        <a href="https://search.google.com/test/rich-results?url={encoded_url}" target="_blank">
                            <button style="padding:8px 16px; border-radius:5px; border:1px solid #ccc; background:rgba(128,128,128,0.1); cursor:pointer;">
                                📡 Test in Google
                            </button>
                        </a>
                        <a href="https://validator.schema.org/#url={encoded_url}" target="_blank">
                            <button style="padding:8px 16px; border-radius:5px; border:1px solid #ccc; background:rgba(128,128,128,0.1); cursor:pointer;">
                                ✅ Schema Validator
                            </button>
                        </a>
                    </div>
                    """, unsafe_allow_html=True)

                with col_json:
                    st.markdown("#### 📝 Raw JSON-LD")
                    st.json(schema_objs, expanded=True)

            else:
                st.info("👆 Click on a row in the table above to see the Google Preview.")
                
                csv = schema_view.to_csv(index=False).encode('utf-8')
                st.download_button("📥 Download Schema Report", csv, "schema_analysis.csv", "text/csv")

        with tab_search:
            st.subheader("🔍 Custom Search Results")
            s_conf = st.session_state.get('search_config', None)
            if 'custom_search_count' not in df.columns:
                st.info("Please configure Custom Search in the Sidebar and run a new crawl.")
            elif s_conf:
                st.info(f"Results for search: **'{s_conf['query']}'** ({s_conf['mode']}) in **{s_conf['scope']}**")
                found_df = df[df['custom_search_count'] > 0][['url', 'custom_search_count']].copy()
                st.dataframe(found_df, use_container_width=True)
            else:
                st.warning("No search was configured.") 
        
        with tab_cannibal:
            st.subheader("👯 Content Similarity & Pruning")
            st.markdown("""
            **Two-Step Pruning Strategy:**
            1. **Duplicates (High Match):** Pages that are copies. Action: *Delete or 301 Redirect.*
            2. **Cannibalization (Medium Match):** Different pages fighting for the same topic. Action: *Merge content.*
            """)
            
            if not HAS_SKLEARN:
                st.error("❌ 'scikit-learn' is not installed. Please run: `pip install scikit-learn`")
            else:
                col1, col2 = st.columns(2)
                
                merge_thresh = col1.slider("Merge Threshold % (Topic Overlap)", 30, 90, 60, help="Finds pages that talk about similar things.")
                dupe_thresh = col2.slider("Duplicate Threshold % (Exact Copies)", 80, 100, 90, help="Finds pages that are almost identical.")
                
                if st.button("🔍 Analyze Content Similarity", type="primary"):
                    with st.spinner("Comparing semantic fingerprints (Title + H1 + Body)..."):
                        if 'scope_content' not in df.columns:
                            st.error("Please re-crawl the website to capture content data.")
                        else:
                            cannibal_df = analyze_content_cannibalization(df, merge_threshold=merge_thresh/100, duplicate_threshold=dupe_thresh/100)
                            st.session_state.cannibal_data = cannibal_df
                
                if 'cannibal_data' in st.session_state:
                    data = st.session_state.cannibal_data
                    
                    if data.empty:
                        st.success("✅ No similarity found above your thresholds.")
                    else:
                        duplicates = data[data['Recommendation'].str.contains("Remove")]
                        st.write("### 🚨 1. Duplicates (High Risk - Remove/Redirect)")
                        if not duplicates.empty:
                            st.error(f"Found {len(duplicates)} pages that look like copies.")
                            st.dataframe(
                                duplicates, 
                                use_container_width=True,
                                column_config={
                                    "Similarity": st.column_config.ProgressColumn("Match %", format="%.1f%%", min_value=0, max_value=100),
                                    "Page A": st.column_config.LinkColumn("Page A"),
                                    "Page B": st.column_config.LinkColumn("Page B"),
                                }
                            )
                        else:
                            st.info("No exact duplicates found.")

                        st.divider()

                        mergers = data[data['Recommendation'].str.contains("Merge")]
                        st.write("### 🤝 2. Merge Opportunities (Keyword Cannibalization)")
                        if not mergers.empty:
                            st.warning(f"Found {len(mergers)} pages covering similar topics.")
                            st.caption("These pages are likely competing for the same keywords. Pick the strongest one and merge the others into it.")
                            st.dataframe(
                                mergers, 
                                use_container_width=True,
                                column_config={
                                    "Similarity": st.column_config.ProgressColumn("Match %", format="%.1f%%", min_value=0, max_value=100),
                                    "Page A": st.column_config.LinkColumn("Page A"),
                                    "Page B": st.column_config.LinkColumn("Page B"),
                                }
                            )
                        else:
                            st.info("No merge opportunities found with current settings.")

        with tab_gsc:
            st.subheader("📈 Google Search Console Analysis")
            
            if not HAS_GSC:
                st.error("Missing libraries. Run: `pip install google-api-python-client google-auth`")
            elif not st.session_state.gsc_service or not gsc_property:
                st.info("👈 Please upload your JSON Key and select a property in the sidebar.")
            else:
                st.write("### 1. Performance Metrics")
                
                valid_dates = False
                start_d, end_d = None, None
                
                if 'date_range' in locals() and isinstance(date_range, tuple) and len(date_range) == 2:
                    start_d, end_d = date_range
                    st.caption(f"Fetching data from: **{start_d}** to **{end_d}**")
                    valid_dates = True
                elif 'date_range' not in locals():
                     st.warning("⚠️ Date picker loading...")
                else:
                    st.warning("⚠️ Please select both a Start Date and an End Date in the sidebar.")

                if st.button("🔄 Fetch Performance (Clicks, Imp, CTR, Pos)", disabled=not valid_dates):
                    with st.spinner("Fetching performance data..."):
                        gsc_data = fetch_gsc_data(st.session_state.gsc_service, gsc_property, start_d, end_d)
                        
                        if not gsc_data.empty:
                            st.success(f"Fetched data for {len(gsc_data)} URLs!")
                            
                            if 'GSC CTR' in gsc_data.columns:
                                gsc_data['GSC CTR'] = (gsc_data['GSC CTR'] * 100).round(2).astype(str) + '%'
                            if 'GSC Position' in gsc_data.columns:
                                gsc_data['GSC Position'] = gsc_data['GSC Position'].round(1)

                            if df is not None and not df.empty:
                                df['url'] = df['url'].astype(str)
                                gsc_data['url'] = gsc_data['url'].astype(str)
                                
                                merged_gsc = pd.merge(df, gsc_data, on='url', how='left')
                                
                                zero_clicks = merged_gsc[merged_gsc['GSC Clicks'].fillna(0) == 0]
                                st.metric("Pages in List with 0 Clicks", len(zero_clicks))
                                
                                display_df = merged_gsc.copy()
                                complex_cols = ['schema_dump', 'internal_links', 'external_links', 
                                              'images', 'redirect_chain', 'header_structure', 'custom_extraction']
                                for col in complex_cols:
                                    if col in display_df.columns:
                                        display_df[col] = display_df[col].astype(str)
                                
                                target_cols = ['url', 'GSC Clicks', 'GSC Impressions', 'GSC CTR', 'GSC Position']
                                final_cols = [c for c in target_cols if c in display_df.columns]
                                
                                st.dataframe(display_df[final_cols], use_container_width=True)
                                st.session_state.gsc_merged_data = merged_gsc
                            else:
                                st.dataframe(gsc_data, use_container_width=True)
                                st.session_state.gsc_merged_data = gsc_data
                        else:
                            st.warning("No data found in GSC for this period.")

                st.divider()

                st.write("### 2. Indexing & Canonical Inspection")
                st.info("ℹ️ Checks the live Google Index status. Use 'List Mode' crawl to define your URLs first.")
                
                available_urls = []
                if df is not None and not df.empty:
                    available_urls = df['url'].tolist()
                elif 'gsc_merged_data' in st.session_state and not st.session_state.gsc_merged_data.empty:
                    available_urls = st.session_state.gsc_merged_data['url'].tolist()
                
                if available_urls:
                    st.write(f"**Ready to inspect {len(available_urls)} URLs**")
                    
                    col_i1, col_i2 = st.columns(2)
                    do_inspect = False
                    urls_to_run = []
                    
                    with col_i1:
                        specific_url = st.text_input("Inspect Single URL")
                        if st.button("🔎 Inspect One"):
                            if specific_url:
                                urls_to_run = [specific_url]
                                do_inspect = True
                            else:
                                st.warning("Enter a URL")

                    with col_i2:
                        st.write(f"Batch Inspect ({len(available_urls)} URLs)")
                        if st.button(f"🚀 Inspect All {len(available_urls)} URLs"):
                            urls_to_run = available_urls
                            do_inspect = True
                    
                    if do_inspect:
                        with st.status("Inspecting URLs... (This can take a while)") as status:
                            insp_df = inspect_url_indexing(st.session_state.gsc_service, gsc_property, urls_to_run)
                            status.update(label="Inspection Complete!", state="complete")
                        
                        if not insp_df.empty:
                            if 'Google Canonical' in insp_df.columns and 'User Canonical' in insp_df.columns:
                                mismatches = insp_df[
                                    (insp_df['Google Canonical'] != 'N/A') & 
                                    (insp_df['Google Canonical'] != insp_df['User Canonical'])
                                ]
                                if not mismatches.empty:
                                    st.error(f"⚠️ Found {len(mismatches)} Canonical Mismatches (Google chose differently)")
                                    st.dataframe(mismatches, use_container_width=True)
                                else:
                                    st.success("✅ Google respects all your canonical tags.")

                            st.dataframe(insp_df, use_container_width=True)
                            csv_insp = insp_df.to_csv(index=False).encode('utf-8')
                            st.download_button("📥 Download Inspection Report", csv_insp, "indexing_report.csv", "text/csv")
                else:
                    st.warning("⚠️ No URLs found to inspect. Please run a Crawl (List Mode) or Fetch Performance Data first.")
                    
                st.divider()
                st.write("### 🚀 3. Bulk Indexing API (Force Indexing)")
                st.info("Directly ping Google to crawl/index these URLs immediately. **Limit: 200 per day.** Ensure your Service Account is added as an 'Owner' in GSC.")
                
                if 'last_gsc_key' not in st.session_state or not st.session_state.last_gsc_key:
                    st.warning("⚠️ Please upload your Service Account JSON Key in the sidebar first.")
                else:
                    c_idx1, c_idx2 = st.columns([2, 1])
                    
                    with c_idx1:
                        avail_urls = df['url'].tolist() if (df is not None and not df.empty) else []
                        selected_index_urls = st.multiselect(
                            "Select Crawled URLs to Submit:", 
                            options=avail_urls,
                            help="Select URLs from your current crawl."
                        )
                        
                        manual_index_urls = st.text_area("Or Paste URLs (One per line):", placeholder="https://mysite.com/new-post")
                    
                    with c_idx2:
                        indexing_action = st.radio(
                            "Action Type:", 
                            ["URL_UPDATED", "URL_DELETED"], 
                            help="UPDATED = Add/Refresh page. DELETED = Remove page from Google."
                        )
                    
                    if st.button("🚀 Ping Google Indexing API", type="primary"):
                        final_urls = list(set(selected_index_urls + [u.strip() for u in manual_index_urls.split('\n') if u.strip()]))
                        
                        if not final_urls:
                            st.error("❌ Please provide at least one URL.")
                        elif len(final_urls) > 200:
                            st.error(f"❌ You selected {len(final_urls)} URLs. The default daily limit is 200. Please reduce your list.")
                        else:
                            with st.spinner(f"Pinging Google for {len(final_urls)} URLs..."):
                                idx_results_df = submit_to_indexing_api(
                                    st.session_state.last_gsc_key, 
                                    final_urls, 
                                    indexing_action
                                )
                                
                                st.success("✅ Submission Process Complete!")
                                st.dataframe(
                                    idx_results_df, 
                                    use_container_width=True,
                                    column_config={
                                        "URL": st.column_config.LinkColumn("Target URL"),
                                        "Status": st.column_config.TextColumn("Status", width="small"),
                                        "Details": st.column_config.TextColumn("API Response")
                                    }
                                )
                                
                                csv_idx = idx_results_df.to_csv(index=False).encode('utf-8')
                                st.download_button("📥 Download Submission Report", csv_idx, "indexing_api_report.csv", "text/csv")
        
        with tab_audit:
            st.subheader("📅 AI Content Relevance & Freshness Auditor")
            st.info("Identify 'Stale' or 'Zombie' pages that need to be updated or removed based on current real-world relevance.")
            
            c1, c2 = st.columns(2)
            with c1:
                aud_provider = st.selectbox("AI Provider", ["Gemini", "OpenAI Compatible (Groq/Ollama/OpenAI)"], key="aud_p")
                aud_key = st.text_input("API Key (or Ollama URL if empty)", type="password", key="aud_k")
            with c2:
                aud_model = st.text_input("Model", value="gemini-1.5-flash" if aud_provider=="Gemini" else "llama-3.3-70b-versatile", key="aud_m")
                aud_url = st.text_input("Endpoint (if needed)", value="https://api.groq.com/openai/v1/chat/completions", key="aud_u")

            audit_scope = st.radio("Audit Scope:", ["Selected Pages (Custom)", "Top 10 High Traffic (GSC Needed)", "All Indexable Pages (WARNING: High API Cost)"], horizontal=True)

            targets_to_audit = []
            if audit_scope == "Selected Pages (Custom)":
                selected_urls = st.multiselect("Select URLs to Audit", df['url'].tolist(), default=df['url'].head(3).tolist())
                targets_to_audit = df[df['url'].isin(selected_urls)].to_dict('records')
                
            elif audit_scope == "Top 10 High Traffic (GSC Needed)":
                if 'gsc_merged_data' in st.session_state and not st.session_state.gsc_merged_data.empty:
                    targets_to_audit = st.session_state.gsc_merged_data.sort_values(by='GSC Clicks', ascending=False).head(10).to_dict('records')
                else:
                    st.warning("⚠️ No GSC Data found. Please fetch GSC data in the 'GSC Data' tab first.")
                    
            else:
                targets_to_audit = df[df['indexability'] == 'Indexable'].to_dict('records')
                st.warning(f"⚠️ You are about to send {len(targets_to_audit)} pages to the AI API. This may take a while.")

            if st.button("🚀 Start Content Audit", type="primary"):
                if not aud_key and aud_provider == "Gemini":
                    st.error("⚠️ Please enter an API key to run the audit.")
                elif not targets_to_audit:
                    st.error("⚠️ No target URLs found to audit.")
                else:
                    progress_aud = st.progress(0)
                    status_aud = st.empty()
                    audit_results = []
                    
                    with ThreadPoolExecutor(max_workers=2) as executor:
                        futures = [executor.submit(analyze_content_freshness, t['url'], t.get('title', ''), t.get('scope_content', ''), aud_provider, aud_key, aud_model, aud_url) for t in targets_to_audit]
                        for i, f in enumerate(as_completed(futures)):
                            audit_results.append(f.result())
                            progress_aud.progress((i + 1) / len(targets_to_audit))
                            status_aud.text(f"Audited {i + 1}/{len(targets_to_audit)}")
                    
                    aud_res_df = pd.DataFrame(audit_results)
                    st.session_state.content_audit_data = aud_res_df
                    status_aud.success("✅ Audit Complete!")

            if 'content_audit_data' in st.session_state and not st.session_state.content_audit_data.empty:
                res = st.session_state.content_audit_data
                
                m1, m2, m3, m4 = st.columns(4)
                m1.metric("KEEP", len(res[res['Decision'] == 'KEEP']))
                m2.metric("UPDATE", len(res[res['Decision'] == 'UPDATE']), delta_color="off")
                m3.metric("REMOVE", len(res[res['Decision'] == 'REMOVE']), delta="-"+str(len(res[res['Decision'] == 'REMOVE'])))
                m4.metric("ERRORS", len(res[res['Decision'] == 'ERROR']))
                
                st.dataframe(res, use_container_width=True, column_config={
                    "Decision": st.column_config.SelectboxColumn("Decision", options=["KEEP", "UPDATE", "REMOVE", "UNKNOWN", "ERROR"]),
                    "URL": st.column_config.LinkColumn("Page URL")
                })
                
                csv_aud = res.to_csv(index=False).encode('utf-8')
                st.download_button("📥 Download Audit Report", csv_aud, "content_audit.csv", "text/csv")

            st.divider()
            st.subheader("📝 Grammar & Spelling Auditor")
            st.info("Checks for spelling and grammar errors. Works for any website.")

            if not HAS_LT:
                st.error("❌ `language_tool_python` is missing. Please run: `pip install language_tool_python`")
            else:
                gram_scope = st.radio("Grammar Scope:", ["Selected Pages Only (Use Filter below)", "All Pages (Slow)"], horizontal=True)
                
                if gram_scope == "Selected Pages Only (Use Filter below)":
                    gram_targets = st.multiselect("Select URLs to Check:", df['url'].tolist(), default=df['url'].head(3).tolist())
                else:
                    gram_targets = df['url'].tolist()

                if st.button("🔍 Check Grammar & Spelling"):
                    if not gram_targets:
                        st.warning("No pages selected.")
                    else:
                        progress_gram = st.progress(0)
                        status_gram = st.empty()
                        gram_results = []
                        
                        tool = None
                        use_cloud_fallback = False
                        
                        try:
                            status_gram.text("Initializing LanguageTool...")
                            tool = language_tool_python.LanguageTool('en-US') 
                        except Exception:
                            status_gram.text("⚠️ Java missing. Switching to Cloud API...")
                            use_cloud_fallback = True

                        total_g = len(gram_targets)
                        
                        for i, u in enumerate(gram_targets):
                            row_data = df[df['url'] == u].iloc[0]
                            
                            check_map = {
                                'Title': row_data.get('title', ''),
                                'H1': row_data.get('h1_tags', ''),
                                'Meta Desc': row_data.get('meta_description', ''),
                                'Content': row_data.get('scope_content', '')[:2500] 
                            }
                            
                            for field, text in check_map.items():
                                if text and len(text) > 5:
                                    matches = []
                                    
                                    if tool:
                                        try:
                                            lt_matches = tool.check(text)
                                            for m in lt_matches:
                                                matches.append({
                                                    'replacements': m.replacements[:3],
                                                    'offset': m.offset,
                                                    'length': m.errorLength
                                                })
                                        except: pass
                                    
                                    elif use_cloud_fallback:
                                        matches = check_grammar_cloud(text)
                                        time.sleep(0.3) 
                                    
                                    for m in matches:
                                        start = m['offset']
                                        end = m['offset'] + m['length']
                                        wrong_word = text[start:end]
                                        
                                        if not wrong_word.strip():
                                            continue

                                        gram_results.append({
                                            'Page URL': u,
                                            'Location': field,
                                            'Wrong Spelling': wrong_word,
                                            'Correct Spelling': ", ".join(m['replacements'])
                                        })
                            
                            progress_gram.progress((i + 1) / total_g)
                            status_gram.text(f"Checking {i+1}/{total_g}: {u}")

                        st.session_state.grammar_audit_data = pd.DataFrame(gram_results)
                        status_gram.success("✅ Grammar Check Complete!")

                if 'grammar_audit_data' in st.session_state and not st.session_state.grammar_audit_data.empty:
                    gdf = st.session_state.grammar_audit_data
                    st.write(f"**Found {len(gdf)} potential errors.**")
                    
                    st.dataframe(
                        gdf, 
                        use_container_width=True,
                        column_config={
                            "Page URL": st.column_config.LinkColumn("Page"),
                            "Wrong Spelling": st.column_config.TextColumn("Wrong Spelling", width="small"),
                            "Correct Spelling": st.column_config.TextColumn("Correct Spelling", width="medium"),
                        }
                    )
                    
                    csv_gram = gdf.to_csv(index=False).encode('utf-8')
                    st.download_button("📥 Download Spelling Report", csv_gram, "spelling_report.csv", "text/csv")
                elif 'grammar_audit_data' in st.session_state:
                    st.info("No spelling errors found in the selected pages.")

        with tab_competitor:
            st.subheader("⚔️ The Ultimate Deep SEO Analysis")
            st.info("Analyzes architecture, semantic gaps, CWV, and deep site-wide interlinking.")
            
            c1, c2 = st.columns([1, 1])
            with c1:
                my_url_input = st.text_input("Your Page URL", key="adv_my_url", placeholder="https://mysite.com/page")
                keywords_input = st.text_area("Target Keywords (Comma separated)", key="adv_kws", placeholder="seo tools, python script")
                
                st.write("**Deep Crawl Settings**")
                enable_deep_inlinks = st.checkbox("🕷️ Crawl Entire Website for Inlinks", help="Finds exact number of internal pages pointing to the target. WARNING: Very slow!")
                max_inlink_pages = st.number_input("Max pages to crawl per domain (if enabled)", 50, 5000, 200, step=50)

            with c2:
                competitors_input = st.text_area("Competitor URLs (1 per line)", height=220, key="adv_comps")

            if st.button("🚀 Run The Ultimate Analysis", type="primary"):
                if not my_url_input or not competitors_input or not keywords_input:
                    st.error("⚠️ Fill in all fields.")
                else:
                    comps = [u.strip() for u in competitors_input.split('\n') if u.strip()]
                    keywords = [k.strip().lower() for k in keywords_input.split(',') if k.strip()]
                    all_urls = [my_url_input] + comps
                    
                    progress_bar = st.progress(0)
                    status_text = st.empty()
                    results_data = []
                    
                    crawler_comp = UltraFrogCrawler() 
                    
                    with ThreadPoolExecutor(max_workers=3) as executor:
                        future_to_url = {executor.submit(crawler_comp.extract_page_data, u): u for u in all_urls}
                        
                        for i, future in enumerate(as_completed(future_to_url)):
                            url = future_to_url[future]
                            status_text.text(f"Fetching Advanced Metrics: {url}")
                            
                            try:
                                data = future.result()
                                parsed_url = urlparse(url)
                                domain = parsed_url.netloc.lower()
                                
                                try:
                                    r = requests.get(url, timeout=8, headers={'User-Agent': 'Mozilla/5.0'})
                                    soup = BeautifulSoup(r.content, 'html.parser')
                                    
                                    list_count = len(soup.find_all(['ul', 'ol']))
                                    table_count = len(soup.find_all('table'))
                                    video_count = len(soup.find_all(['iframe', 'video']))
                                    
                                    paragraphs = soup.find_all('p')
                                    p_lengths = [len(p.get_text().split()) for p in paragraphs if len(p.get_text().strip()) > 0]
                                    avg_p_len = sum(p_lengths) / len(p_lengths) if p_lengths else 0
                                    
                                    publish_date = extract_publish_date(soup)
                                except:
                                    list_count, table_count, video_count, avg_p_len, publish_date = 0, 0, 0, 0, "Unknown"
                                
                                title = str(data.get('title', ''))
                                desc = str(data.get('meta_description', ''))
                                h1 = str(data.get('h1_tags', '')).lower()
                                h2 = str(data.get('h2_tags', '')).lower()
                                h3 = str(data.get('h3_tags', '')).lower()
                                body = str(data.get('scope_content', ''))
                                body_lower = body.lower()
                                
                                alts = " ".join([str(img.get('alt', '')).lower() for img in data.get('images', [])])
                                word_count = data.get('word_count', 0)
                                image_count = data.get('image_count', 0)
                                first_100_words = " ".join(body_lower.split()[:100])
                                
                                word_to_img_ratio = round(word_count / image_count, 1) if image_count > 0 else word_count
                                readability = textstat.flesch_reading_ease(body) if body else 0
                                domain_age = get_domain_age(domain)
                                
                                ext_links = data.get('external_links', [])
                                nofollow_count = sum(1 for link in ext_links if 'nofollow' in link.get('rel_status', ''))
                                ext_total = len(ext_links)
                                nofollow_ratio = f"{int((nofollow_count/ext_total)*100)}%" if ext_total > 0 else "0%"
                                
                                total_inlinks_to_page = "Skipped"
                                if enable_deep_inlinks:
                                    status_text.text(f"🕷️ Deep Crawling {domain} to find inlinks...")
                                    total_inlinks_to_page = deep_crawl_for_inlinks(url, max_pages=max_inlink_pages)
                                
                                kw_stats = {}
                                for kw in keywords:
                                    fuzzy_count = count_fuzzy_match(body, kw)
                                    kw_stats[kw] = {
                                        'Domain': domain.count(kw), 'URL': parsed_url.path.lower().count(kw),
                                        'Title': title.lower().count(kw), 'Desc': desc.lower().count(kw),
                                        'H1': h1.count(kw), 'H2': h2.count(kw), 'H3': h3.count(kw),
                                        'Alt': alts.count(kw), 'Body Exact': body_lower.count(kw),
                                        'Body Fuzzy': fuzzy_count,
                                        'In First 100': "✅" if kw in first_100_words else "❌"
                                    }

                                lcp, cls, mobile_score = "N/A", "N/A", "N/A"
                                if psi_api_key:
                                    status_text.text(f"Running PSI: {url}")
                                    m_res = run_psi_test(url, psi_api_key, "mobile")
                                    if isinstance(m_res, dict) and "Score" in m_res:
                                        mobile_score = m_res['Score']
                                        lcp = m_res.get('LCP', 'N/A')
                                        cls = m_res.get('CLS', 'N/A')

                                results_data.append({
                                    'Type': '🟦 ME' if url == my_url_input else '🟥 COMP',
                                    'URL': url,
                                    'Date Published': publish_date,
                                    'Domain Age': domain_age,
                                    'Readability': round(readability, 1),
                                    'Word Count': word_count,
                                    'Avg Paragraph': f"{round(avg_p_len)} words",
                                    'Title / Meta Length': f"{len(title)} / {len(desc)}",
                                    'H2 / H3 Count': f"{data.get('h2_count', 0)} / {data.get('h3_count', 0)}",
                                    'Word/Img Ratio': f"{word_to_img_ratio}:1",
                                    'Media (Vid/List/Tab)': f"{video_count} / {list_count} / {table_count}",
                                    'On-Page Internal Links': data.get('internal_links_count', 0),
                                    'Total Inlinks (Whole Site)': total_inlinks_to_page, 
                                    'External Outbound': f"{ext_total} (Nofollow: {nofollow_ratio})",
                                    'Schema': data.get('schema_types', 'None'),
                                    'CWV (Score | LCP | CLS)': f"{mobile_score} | {lcp} | {cls}",
                                    'Keywords': kw_stats,
                                    'RawBody': body
                                })
                                
                            except Exception as e:
                                st.error(f"Failed {url}: {e}")
                                
                            progress_bar.progress((i + 1) / len(all_urls))
                    
                    status_text.success("✅ Deep Analysis Complete!")
                    
                    my_data = next((item for item in results_data if 'ME' in item['Type']), None)
                    comp_data = [item for item in results_data if 'COMP' in item['Type']]
                    
                    def get_clean_domain(url_str):
                        return urlparse(url_str).netloc.replace('www.', '')

                    st.write("### 🏗️ 1. Content & Technical Architecture")
                    st.markdown("Side-by-side structural comparison. **Metrics are rows, Websites are columns.**")
                    
                    if my_data:
                        metrics_keys = [
                            'Word Count', 'Readability', 'H2 / H3 Count', 'Images', 
                            'Word/Img Ratio', 'Media (Vid/List/Tab)', 'On-Page Internal Links', 
                            'Total Inlinks (Whole Site)', 'External Outbound', 'Schema', 
                            'Domain Age', 'Date Published', 'CWV (Score | LCP | CLS)'
                        ]
                        
                        t1_rows = []
                        for metric in metrics_keys:
                            row = {"Metric": metric}
                            row["🫵 My Page"] = my_data.get(metric, "N/A")
                            
                            for i, comp in enumerate(comp_data):
                                col_name = f"Comp {i+1} ({get_clean_domain(comp['URL'])})"
                                row[col_name] = comp.get(metric, "N/A")
                                
                            t1_rows.append(row)
                            
                        struct_df = pd.DataFrame(t1_rows)
                        st.dataframe(struct_df, use_container_width=True)
                    else:
                        st.error("Could not fetch data for 'Your Page URL'.")

                    st.divider()

                    st.write("### 🎯 2. Keyword Zone Matrix")
                    st.info("Select a keyword below to see exactly which HTML tags your competitors placed it in.")
                    
                    if my_data:
                        kw_tabs = st.tabs(keywords)
                        flat_kw_export = [] 
                        
                        for i, kw in enumerate(keywords):
                            with kw_tabs[i]:
                                zones = ['Body Exact', 'Body Fuzzy', 'In First 100 Words?', 'Title', 'H1', 'H2', 'H3', 'Meta Desc', 'Img Alt', 'URL', 'Domain']
                                kw_rows = []
                                
                                for zone in zones:
                                    row = {"HTML Zone": zone}
                                    my_val = my_data['Keywords'][kw].get(zone, 0)
                                    row["🫵 My Page"] = my_val
                                    
                                    comp_vals_for_avg = []
                                    for j, comp in enumerate(comp_data):
                                        comp_val = comp['Keywords'][kw].get(zone, 0)
                                        col_name = f"Comp {j+1} ({get_clean_domain(comp['URL'])})"
                                        row[col_name] = comp_val
                                        
                                        if isinstance(comp_val, (int, float)):
                                            comp_vals_for_avg.append(comp_val)
                                    
                                    if comp_vals_for_avg:
                                        avg = sum(comp_vals_for_avg) / len(comp_vals_for_avg)
                                        row["Competitor Avg"] = round(avg, 1)
                                        
                                        if isinstance(my_val, (int, float)):
                                            if my_val < avg: row["Status"] = "📉 Below Avg"
                                            elif my_val > avg * 2: row["Status"] = "📈 High (Careful)"
                                            else: row["Status"] = "✅ Optimal"
                                    else:
                                        row["Competitor Avg"] = "N/A"
                                        row["Status"] = "-"
                                        
                                    kw_rows.append(row)
                                    
                                    flat_row = {"Keyword": kw, **row}
                                    flat_kw_export.append(flat_row)
                                
                                kw_df = pd.DataFrame(kw_rows)
                                st.dataframe(kw_df, use_container_width=True)

                    st.divider()

                    st.write("### 🧠 3. Semantic LSI Gap Analysis")
                    st.markdown("Top 2-word and 3-word phrases used by competitors. See exactly who is using which sub-topics.")
                    
                    try:
                        comp_data = [r for r in results_data if 'COMP' in r['Type']]
                        comp_bodies = [r['RawBody'] for r in comp_data]
                        my_body = my_data['RawBody'].lower() if my_data else ""
                        
                        if comp_bodies:
                            vec = CountVectorizer(ngram_range=(2, 3), stop_words='english', max_features=30)
                            matrix = vec.fit_transform(comp_bodies)
                            sums = matrix.sum(axis=0).A1
                            vocab = vec.get_feature_names_out()
                            
                            freq = [(vocab[idx], sums[idx]) for idx in range(len(vocab))]
                            freq = sorted(freq, key=lambda x: x[1], reverse=True)
                            
                            lsi_data = []
                            for phrase, total_count in freq[:20]:
                                my_usage_count = my_body.count(phrase)
                                status = "✅ Used" if my_usage_count > 0 else "❌ Missing"
                                
                                row_data = {
                                    "Semantic Phrase (LSI)": phrase,
                                    "🫵 My Page": my_usage_count,
                                }
                                
                                for i, comp in enumerate(comp_data):
                                    col_name = f"Comp {i+1} ({get_clean_domain(comp['URL'])})"
                                    row_data[col_name] = comp['RawBody'].lower().count(phrase)
                                    
                                row_data["Gap Status"] = status
                                lsi_data.append(row_data)
                                
                            lsi_df = pd.DataFrame(lsi_data)
                            st.dataframe(
                                lsi_df, 
                                use_container_width=True,
                                column_config={
                                    "Gap Status": st.column_config.TextColumn("Status", width="small")
                                }
                            )
                        else:
                            st.warning("Not enough competitor data to generate LSI keywords.")
                    except Exception as e:
                        st.error(f"Could not generate LSI phrases: {e}")

                    st.markdown("<br>", unsafe_allow_html=True)
                    c_btn1, c_btn2, c_btn3 = st.columns(3)
                    
                    if 'struct_df' in locals():
                        c_btn1.download_button("📥 Download Architecture", struct_df.to_csv(index=False).encode('utf-8'), "competitor_structure.csv", "text/csv")
                    if 'flat_kw_export' in locals() and flat_kw_export:
                        flat_df = pd.DataFrame(flat_kw_export)
                        c_btn2.download_button("📥 Download Keyword Matrix", flat_df.to_csv(index=False).encode('utf-8'), "competitor_keywords.csv", "text/csv")
                    if 'lsi_df' in locals():
                        c_btn3.download_button("📥 Download LSI Gaps", lsi_df.to_csv(index=False).encode('utf-8'), "lsi_gaps.csv", "text/csv")
        
        st.divider()
        st.header("📥 Full Report")
        
        if st.session_state.storage_mode == "SQLite" and os.path.exists(st.session_state.db_file):
            with open(st.session_state.db_file, "rb") as file:
                st.download_button("📊 Download Complete Database (SQLite)", file, "battersea_data.db", "application/octet-stream", type="primary")
        else:
            csv_all = df.to_csv(index=False).encode('utf-8')
            st.download_button("📊 Download Complete Crawl Data", csv_all, "complete_crawl.csv", "text/csv", type="primary")

    else:
        st.info("👈 Configure your crawl settings in the sidebar and click '🚀 Start Crawl' to begin.")
        
# ==========================================
# 3. SOCIAL MEDIA WORKSPACE
# ==========================================
elif workspace == "📱 Social Media Reporting":
    
    # --- API HELPER FUNCTIONS ---
    def fetch_meta_data(page_id, ig_id, token, start_date, end_date):
        """Fetches REAL data with Deep Fetch for Post-Level Impressions filtered by Date"""
        if not token: return None, None, []
        
        fb_data, ig_data = {}, {}
        all_posts = [] 
        
        # ---> THIS IS THE NEW DATE FORMATTING LOGIC <---
        # We must convert Python dates into "YYYY-MM-DD" text so Facebook understands it
        since_str = start_date.strftime('%Y-%m-%d')
        until_str = (end_date + timedelta(days=1)).strftime('%Y-%m-%d')
        
        try:
            # ==========================================
            # 1. FACEBOOK PAGE DATA
            # ==========================================
            page_token = token 
            
            if page_id:
                # 1A: Get Page Access Token
                token_url = f"https://graph.facebook.com/v19.0/{page_id}?fields=access_token&access_token={token}"
                token_res = requests.get(token_url).json()
                page_token = token_res.get('access_token', token)
                
                # 1B: Get Page Followers
                fb_info_url = f"https://graph.facebook.com/v19.0/{page_id}?fields=followers_count,fan_count&access_token={page_token}"
                fb_info_res = requests.get(fb_info_url).json()
                followers = fb_info_res.get('followers_count', fb_info_res.get('fan_count', 0))
                
                # 1C: Get Page Impressions
                impressions = 0
                fb_insights_url = f"https://graph.facebook.com/v19.0/{page_id}/insights?metric=page_impressions&period=day&access_token={page_token}"
                fb_insights_res = requests.get(fb_insights_url).json()
                if "data" in fb_insights_res and len(fb_insights_res["data"]) > 0:
                    vals = fb_insights_res["data"][0].get("values", [])
                    impressions = sum(v.get("value", 0) for v in vals)
                
                # 1D: Get Posts & DEEP FETCH Impressions (NOW USING SINCE/UNTIL DATES)
                engagements = 0
                fb_posts_url = f"https://graph.facebook.com/v19.0/{page_id}/published_posts?fields=id,created_time,permalink_url,likes.summary(true),comments.summary(true),shares&since={since_str}&until={until_str}&limit=100&access_token={page_token}"
                fb_posts_res = requests.get(fb_posts_url).json()
                
                if "data" in fb_posts_res:
                    for post in fb_posts_res["data"]:
                        post_id = post.get("id")
                        likes = post.get("likes", {}).get("summary", {}).get("total_count", 0)
                        comments = post.get("comments", {}).get("summary", {}).get("total_count", 0)
                        shares = post.get("shares", {}).get("count", 0)
                        engagements += (likes + comments + shares)
                        
                        # --- THE DEEP FETCH FOR FB IMPRESSIONS ---
                        post_imps = "-" # Default to a clean dash
                        try:
                            p_ins_url = f"https://graph.facebook.com/v19.0/{post_id}/insights?metric=post_impressions&access_token={page_token}"
                            p_ins_res = requests.get(p_ins_url).json()
                            
                            # If Facebook returns an error (like on a Reel), keep it a clean dash
                            if "error" in p_ins_res:
                                post_imps = "-" 
                            # If it succeeds, grab the actual number
                            elif "data" in p_ins_res and len(p_ins_res["data"]) > 0:
                                post_imps = p_ins_res["data"][0]["values"][0]["value"]
                        except: 
                            pass
                        
                        date_str = post.get("created_time", "").split("T")[0] if "created_time" in post else "Unknown"
                        all_posts.append({
                            "Date": date_str,
                            "Platform": "Facebook",
                            "Post Type": "Post",
                            "Preview Link": post.get("permalink_url", "https://facebook.com"),
                            "Impressions": post_imps,
                            "Likes": likes,
                            "Comments": comments,
                            "Shares/Reposts": shares
                        })
                
                er = round((engagements / impressions) * 100, 2) if impressions > 0 else 0
                
                fb_data = {
                    "Followers": followers, 
                    "Impressions": impressions, 
                    "Engagements": engagements,
                    "ER %": f"{er}%"
                }

            # ==========================================
            # 1.5. AUTO-DISCOVER INSTAGRAM ID
            # ==========================================
            actual_ig_id = ig_id 
            if page_id:
                find_ig_url = f"https://graph.facebook.com/v19.0/{page_id}?fields=instagram_business_account&access_token={page_token}"
                find_ig_res = requests.get(find_ig_url).json()
                if 'instagram_business_account' in find_ig_res:
                    actual_ig_id = find_ig_res['instagram_business_account']['id']
            
            # ==========================================
            # 2. INSTAGRAM ACCOUNT DATA
            # ==========================================
            if actual_ig_id:
                clean_ig_id = str(actual_ig_id).replace("@", "").strip()
                
                # Get IG Followers
                ig_info_url = f"https://graph.facebook.com/v19.0/{clean_ig_id}?fields=followers_count&access_token={token}"
                ig_info_res = requests.get(ig_info_url).json()
                ig_followers = ig_info_res.get('followers_count', 0) if "error" not in ig_info_res else 0
                    
                # Get IG Overall Reach
                ig_impressions = 0
                ig_insights_url = f"https://graph.facebook.com/v19.0/{clean_ig_id}/insights?metric=reach&period=day&access_token={token}"
                ig_insights_res = requests.get(ig_insights_url).json()
                if "data" in ig_insights_res and len(ig_insights_res["data"]) > 0:
                    vals = ig_insights_res["data"][0].get("values", [])
                    ig_impressions = sum(v.get("value", 0) for v in vals)
                
                # Get IG Media & DEEP FETCH Reach (NOW USING SINCE/UNTIL DATES)
                ig_engagements = 0
                ig_media_url = f"https://graph.facebook.com/v19.0/{clean_ig_id}/media?fields=id,timestamp,permalink,media_type,like_count,comments_count&since={since_str}&until={until_str}&limit=100&access_token={token}"
                ig_media_res = requests.get(ig_media_url).json()
                
                if "data" in ig_media_res:
                    for media in ig_media_res["data"]:
                        m_id = media.get("id")
                        likes = media.get("like_count", 0)
                        comments = media.get("comments_count", 0)
                        ig_engagements += (likes + comments)
                        
                        # --- THE DEEP FETCH FOR IG REACH ---
                        ig_post_reach = "N/A"
                        try:
                            m_ins_url = f"https://graph.facebook.com/v19.0/{m_id}/insights?metric=reach&access_token={token}"
                            m_ins_res = requests.get(m_ins_url).json()
                            if "data" in m_ins_res and len(m_ins_res["data"]) > 0:
                                ig_post_reach = m_ins_res["data"][0]["values"][0]["value"]
                        except: pass
                        
                        date_str = media.get("timestamp", "").split("T")[0] if "timestamp" in media else "Unknown"
                        all_posts.append({
                            "Date": date_str,
                            "Platform": "Instagram",
                            "Post Type": media.get("media_type", "Media"),
                            "Preview Link": media.get("permalink", "https://instagram.com"),
                            "Impressions": ig_post_reach,
                            "Likes": likes,
                            "Comments": comments,
                            "Shares/Reposts": "API Restricted" 
                        })
                
                ig_er = round((ig_engagements / ig_impressions) * 100, 2) if ig_impressions > 0 else 0
                            
                ig_data = {
                    "Followers": ig_followers, 
                    "Impressions": ig_impressions, 
                    "Engagements": ig_engagements, 
                    "ER %": f"{ig_er}%"
                }
                    
            return fb_data if fb_data else None, ig_data if ig_data else None, all_posts
            
        except Exception as e:
            st.error(f"🚨 Python Request Error: {e}")
            return None, None, []

    def fetch_linkedin_data(org_urn, token):
        """Fetches data from LinkedIn API"""
        if not token or not org_urn: return None
        headers = {
            "Authorization": f"Bearer {token}",
            "X-Restli-Protocol-Version": "2.0.0",
            "LinkedIn-Version": "202401"
        }
        try:
            # Example endpoint for page statistics
            li_url = f"https://api.linkedin.com/rest/organizationalEntityPageStatistics?q=organization&organization={org_urn}"
            li_res = requests.get(li_url, headers=headers).json()
            return {"Followers": 0, "Impressions": 0, "Engagements": 0} # Parse real response here
        except Exception as e:
            st.error(f"LinkedIn API Error: {e}")
            return None

    def fetch_x_data(username, token):
        """Fetches data from X (Twitter) API v2"""
        if not token or not username: return None
        headers = {"Authorization": f"Bearer {token}"}
        try:
            # First get User ID from username
            user_url = f"https://api.twitter.com/2/users/by/username/{username}?user.fields=public_metrics"
            user_res = requests.get(user_url, headers=headers).json()
            if "data" in user_res:
                metrics = user_res["data"]["public_metrics"]
                return {"Followers": metrics.get("followers_count", 0), "Impressions": 0, "Engagements": 0}
            return None
        except Exception as e:
            st.error(f"X API Error: {e}")
            return None

    # --- CREDENTIALS MANAGEMENT (MULTI-CLIENT) ---
    SMO_CRED_FILE = "smo_credentials.json"
    
    def load_smo_creds():
        if os.path.exists(SMO_CRED_FILE):
            with open(SMO_CRED_FILE, "r") as f:
                return json.load(f)
        return {} # Returns an empty dict if no file exists

    def save_smo_creds(creds_dict):
        with open(SMO_CRED_FILE, "w") as f:
            json.dump(creds_dict, f, indent=4)

    saved_creds = load_smo_creds()

    # --- SIDEBAR CONFIGURATION ---
    with st.sidebar:
        st.markdown("### 👤 Client Details")
        
        # 1. Profile Selector
        client_options = list(saved_creds.keys())
        client_options.insert(0, "➕ Create New Client...")
        
        selected_profile = st.selectbox("Select Client Profile", client_options)
        
        # 2. Setup variables based on selection
        if selected_profile == "➕ Create New Client...":
            client_name = st.text_input("New Client Name", placeholder="e.g., Travio World")
            c_data = {} # Empty data for new client
        else:
            client_name = selected_profile
            c_data = saved_creds.get(selected_profile, {})
            # Read-only display of the selected name to avoid confusion
            st.text_input("Editing Client:", value=client_name, disabled=True) 

        # 3. Date Range
        today = datetime.now()
        thirty_days_ago = today - timedelta(days=30)
        date_range_smo = st.date_input("Reporting Period", value=(thirty_days_ago, today), format="DD/MM/YYYY")
        
        # 4. IDs and Tokens (Auto-filled if c_data exists)
        with st.expander("🔗 Platform Page IDs", expanded=(selected_profile == "➕ Create New Client...")):
            fb_page_id = st.text_input("Facebook Page ID", value=c_data.get("fb_page_id", ""))
            ig_account_id = st.text_input("Instagram Account ID", value=c_data.get("ig_account_id", ""))
            li_org_urn = st.text_input("LinkedIn Org URN", value=c_data.get("li_org_urn", ""))
            x_username = st.text_input("X (Twitter) Username", value=c_data.get("x_username", ""))
            
        with st.expander("🔑 API Authentication", expanded=False):
            meta_token = st.text_input("Meta Access Token", type="password", value=c_data.get("meta_token", ""))
            li_token = st.text_input("LinkedIn Bearer Token", type="password", value=c_data.get("li_token", ""))
            x_token = st.text_input("X Bearer Token", type="password", value=c_data.get("x_token", ""))
            
        # 5. Save Button
        if st.button("💾 Save Client Profile"):
            if not client_name:
                st.error("Please enter a Client Name before saving.")
            else:
                # Update the specific client's data in the master dictionary
                saved_creds[client_name] = {
                    "fb_page_id": fb_page_id,
                    "ig_account_id": ig_account_id,
                    "li_org_urn": li_org_urn,
                    "x_username": x_username,
                    "meta_token": meta_token,
                    "li_token": li_token,
                    "x_token": x_token
                }
                save_smo_creds(saved_creds)
                st.success(f"Saved profile for {client_name}!")
                time.sleep(1)
                st.rerun()
                
        # Optional: Delete Profile Button
        if selected_profile != "➕ Create New Client...":
            if st.button("🗑️ Delete Profile"):
                del saved_creds[selected_profile]
                save_smo_creds(saved_creds)
                st.warning(f"Deleted profile: {selected_profile}")
                time.sleep(1)
                st.rerun()
            
        st.divider()
        fetch_smo_data = st.button("🔄 Fetch Live Data", type="primary", use_container_width=True)

    # --- INITIALIZE SESSION STATE FOR SOCIAL DATA ---
    if 'social_overview_df' not in st.session_state:
        st.session_state.social_overview_df = pd.DataFrame(columns=["Platform", "Followers", "Impressions", "Engagements", "ER %"])
    if 'social_posts_df' not in st.session_state:
        st.session_state.social_posts_df = pd.DataFrame(columns=["Date", "Platform", "Post Type", "Preview Link", "Impressions", "Likes", "Comments"])

    # --- MAIN DASHBOARD ---
    st.header(f"📱 Social Media Reporting: {client_name if client_name else 'Select Client'}")
    
    # Tabs
    tab_overview, tab_platforms, tab_posts, tab_ai_writer, tab_reports = st.tabs([
        "📊 Executive Overview", "📱 Platform Deep-Dive", "📝 Top Posts", "🤖 AI Content Hub", "📥 Client Reports"
    ])
    
    # --- FETCH DATA LOGIC ---
    if fetch_smo_data:
        if len(date_range_smo) != 2:
            st.error("Please select both a Start Date and End Date in the sidebar.")
        else:
            with st.spinner("Connecting to platform APIs..."):
                start_d, end_d = date_range_smo[0], date_range_smo[1]
                
                # 1. Fetch Meta (Now passing start_d and end_d)
                fb_stats, ig_stats, meta_posts = fetch_meta_data(fb_page_id, ig_account_id, meta_token, start_d, end_d)
            
            # 2. Fetch LinkedIn
            li_stats = fetch_linkedin_data(li_org_urn, li_token)
            # 3. Fetch X
            x_stats = fetch_x_data(x_username, x_token)
            
            # Combine into list
            api_results = []
            if fb_stats: api_results.append({"Platform": "Facebook", **fb_stats})
            if ig_stats: api_results.append({"Platform": "Instagram", **ig_stats})
            if li_stats: api_results.append({"Platform": "LinkedIn", **li_stats})
            if x_stats: api_results.append({"Platform": "X (Twitter)", **x_stats})
            
            if api_results:
                st.session_state.social_overview_df = pd.DataFrame(api_results)
                
                # SAVING POSTS FOR TAB 3!
                if meta_posts:
                    st.session_state.social_posts_df = pd.DataFrame(meta_posts)
                
                st.toast("Live API Data successfully loaded!", icon="✅")
            else:
                st.warning("No API data fetched. Check your tokens and IDs.")
    # -------------------------------------------------------------
    # TAB 1: EXECUTIVE OVERVIEW
    # -------------------------------------------------------------
    with tab_overview:
        st.subheader("📊 Cross-Platform Performance")
        
        df_overview = st.session_state.social_overview_df
        
        if df_overview.empty:
            st.info("👈 Enter your client IDs and API tokens in the sidebar, then click 'Fetch Live Data'.")
        else:
            # Calculate Totals (Using .columns for best practice)
            tot_followers = df_overview['Followers'].sum() if 'Followers' in df_overview.columns else 0
            tot_imps = df_overview['Impressions'].sum() if 'Impressions' in df_overview.columns else 0
            tot_eng = df_overview['Engagements'].sum() if 'Engagements' in df_overview.columns else 0
            
            # Calculate Total Average Engagement Rate
            if tot_imps > 0:
                avg_er = round((tot_eng / tot_imps) * 100, 2)
                avg_er_display = f"{avg_er}%"
            else:
                avg_er_display = "0%"

            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Total Audience", f"{tot_followers:,}")
            m2.metric("Total Impressions", f"{tot_imps:,}")
            m3.metric("Total Engagements", f"{tot_eng:,}")
            m4.metric("Avg. Engagement Rate", avg_er_display) 
            
            st.divider()
            st.write("#### Platform Breakdown")
            st.dataframe(df_overview, use_container_width=True)
        
    # -------------------------------------------------------------
    # TAB 2: PLATFORM DEEP DIVE
    # -------------------------------------------------------------
    with tab_platforms:
        st.subheader("📱 Platform Deep-Dive")
        selected_platform = st.selectbox("Select Platform to Inspect:", ["Instagram", "Facebook", "LinkedIn", "X (Twitter)"])
        st.info("In the future, map specific API demographic and geography arrays to this table.")
        st.dataframe(pd.DataFrame(columns=["Audience Segment", "Value"]), use_container_width=True)
            
    # -------------------------------------------------------------
    # TAB 3: TOP POSTS
    # -------------------------------------------------------------
    with tab_posts:
        st.subheader("📝 Top Performing Posts")
        
        df_posts = st.session_state.social_posts_df
        
        if df_posts.empty:
             st.info("👈 Click 'Fetch Live Data' in the sidebar to populate your posts.")
        else:
            # 1. Filter DataFrames by Platform
            fb_posts = df_posts[df_posts['Platform'] == 'Facebook'].copy()
            ig_posts = df_posts[df_posts['Platform'] == 'Instagram'].copy()
            
            # Common Column Config (Turns the URL into a clickable link)
            col_conf = {
                "Preview Link": st.column_config.LinkColumn("View Post")
            }
            
            # 2. Render Facebook Table
            st.markdown("### 🟦 Facebook Posts")
            if not fb_posts.empty:
                # We drop the 'Platform' column here because it's redundant now
                st.dataframe(
                    fb_posts.drop(columns=['Platform']), 
                    use_container_width=True, 
                    column_config=col_conf
                )
            else:
                st.info("No Facebook posts found for this period.")
                
            st.divider()
            
            # 3. Render Instagram Table
            st.markdown("### 📸 Instagram Posts")
            if not ig_posts.empty:
                st.dataframe(
                    ig_posts.drop(columns=['Platform']), 
                    use_container_width=True, 
                    column_config=col_conf
                )
            else:
                st.info("No Instagram posts found for this period.")
        
    # -------------------------------------------------------------
    # TAB 4: AI CONTENT HUB
    # -------------------------------------------------------------
    with tab_ai_writer:
        st.subheader("🤖 AI Content Hub")
        st.info("Generate social media captions instantly using your SEO AI models.")
        
        c_topic, c_tone = st.columns([3, 1])
        topic = c_topic.text_area("What is the post about? (Paste URL or text)", height=100)
        tone = c_tone.selectbox("Tone of Voice", ["Professional", "Witty/Humorous", "Educational", "Urgent/Sales"])
        
        if st.button("✨ Generate Captions", type="primary"):
            if topic:
                with st.spinner("Writing captions..."):
                    # REUSE YOUR EXISTING AI FUNCTION HERE
                    prompt = f"Write a social media post about this topic: '{topic}'. Tone: {tone}. Give me one version for LinkedIn and one for X (Twitter)."
                    
                    # Assuming you have the Groq/Gemini variables loaded from the SEO tab, 
                    # or you can add them to the Social sidebar. For now, simulating output:
                    time.sleep(1.5)
                    st.success("Feature ready! Connect your Groq/Gemini function here.")
            else:
                st.warning("Please enter a topic first.")
                
    # -------------------------------------------------------------
    # TAB 5: CLIENT REPORTS
    # -------------------------------------------------------------
    with tab_reports:
        st.subheader("📥 Client Reporting")
        st.write(f"Generate the end-of-month report for **{client_name if client_name else 'the client'}**.")
        
        if st.button("🧠 Generate Executive Summary (AI)"):
            if st.session_state.social_overview_df.empty:
                st.warning("Fetch API data first so the AI has numbers to analyze.")
            else:
                st.info("Pass `st.session_state.social_overview_df` to your AI prompt to write a summary.")
        
        st.divider()
        
        if not st.session_state.social_overview_df.empty:
            csv_export = st.session_state.social_overview_df.to_csv(index=False).encode('utf-8')
            file_name = f"{client_name.replace(' ', '_')}_Social_Report.csv" if client_name else "Social_Report.csv"
            
            st.download_button(
                label="📊 Download Complete CSV Report",
                data=csv_export,
                file_name=file_name,
                mime="text/csv",
                type="primary"
            )

# ==========================================
# 4. CONTENT WRITERS WORKSPACE
# ==========================================
elif workspace == "✍️ Content Writers Hub":
    st.header("✍️ Content Writers Hub")
    st.info("Add your AI Generation, Grammar checking, and Plagiarism tools here.")
    # You can move your LanguageTool logic here if you want it decoupled from the crawler

# ==========================================
# 5. OTHER TOOLS (UTILITIES) WORKSPACE
# ==========================================
elif workspace == "🛠️ Other Tools":
    st.header("🛠️ Agency Utility Tools")
    st.markdown("A collection of formatting, compression, and conversion tools for the team.")

    tab_img, tab_pdf, tab_html = st.tabs(["🖼️ Image Tools", "📄 PDF Tools", "🌐 HTML & Doc Tools"])

    # ---------------------------------------------------------
    # TAB 1: IMAGE TOOLS
    # ---------------------------------------------------------
    with tab_img:
        st.subheader("🗜️ Image Optimizer (Bulk & Interactive)")
        st.info("Resize and compress heavy images into SEO-friendly WebP format.")
        
        up_imgs = st.file_uploader("Upload Images for Optimization", type=['png', 'jpg', 'jpeg'], accept_multiple_files=True, key="opt_imgs")
        
        interactive_mode = st.checkbox("🔍 Interactive Preview Mode (Review & adjust images one by one)", value=False)
        st.divider()

        # --- STATE MANAGEMENT FOR INTERACTIVE MODE ---
        if 'int_img_index' not in st.session_state:
            st.session_state.int_img_index = 0
        if 'int_processed_imgs' not in st.session_state:
            st.session_state.int_processed_imgs = {}
        if 'last_upload_names' not in st.session_state:
            st.session_state.last_upload_names = []

        # Reset the wizard if the user uploads a new batch of files
        current_names = [f.name for f in up_imgs] if up_imgs else []
        if current_names != st.session_state.last_upload_names:
            st.session_state.last_upload_names = current_names
            st.session_state.int_img_index = 0
            st.session_state.int_processed_imgs = {}

        if not up_imgs:
            st.write("Waiting for uploads...")
            
        elif not interactive_mode:
            # ==========================================
            # STANDARD BULK MODE (Original Logic)
            # ==========================================
            c_opt1, c_opt2 = st.columns(2)
            max_width = c_opt1.number_input("Max Width (px)", value=1200, step=100)
            quality = c_opt2.slider("WebP Quality (1-100)", 10, 100, 75)
            
            if st.button("🚀 Compress All to WebP", type="primary"):
                with st.spinner("Compressing images..."):
                    zip_buf = io.BytesIO()
                    with zipfile.ZipFile(zip_buf, "a", zipfile.ZIP_DEFLATED, False) as zf:
                        for img_file in up_imgs:
                            try:
                                img = Image.open(img_file)
                                if img.mode in ("RGBA", "P"): img = img.convert("RGB")
                                if img.width > max_width:
                                    ratio = max_width / float(img.width)
                                    new_h = int((float(img.height) * float(ratio)))
                                    img = img.resize((max_width, new_h), Image.Resampling.LANCZOS)
                                    
                                img_byte_arr = io.BytesIO()
                                img.save(img_byte_arr, format='WEBP', quality=quality, optimize=True)
                                
                                orig_name = img_file.name.rsplit('.', 1)[0]
                                zf.writestr(f"{orig_name}_optimized.webp", img_byte_arr.getvalue())
                            except Exception as e: st.error(f"Failed {img_file.name}: {e}")
                                
                    st.success("✅ Bulk Compression Complete!")
                    st.download_button("📥 Download Optimized Images (ZIP)", zip_buf.getvalue(), "optimized_images.zip", "application/zip")

        else:
            # ==========================================
            # NEW INTERACTIVE PREVIEW MODE
            # ==========================================
            if st.session_state.int_img_index < len(up_imgs):
                current_idx = st.session_state.int_img_index
                current_file = up_imgs[current_idx]
                
                st.write(f"### Image {current_idx + 1} of {len(up_imgs)}: `{current_file.name}`")
                
                try:
                    # 1. Load Original Data
                    orig_bytes = current_file.getvalue()
                    orig_img = Image.open(io.BytesIO(orig_bytes))
                    orig_kb = len(orig_bytes) / 1024
                    
                    # 2. Interactive Controls (Keys must be unique per image so sliders reset)
                    c_ctrl1, c_ctrl2 = st.columns(2)
                    dyn_width = c_ctrl1.number_input("Max Width (px)", value=min(1200, orig_img.width), step=100, key=f"w_{current_idx}")
                    dyn_quality = c_ctrl2.slider("WebP Quality %", 10, 100, 75, key=f"q_{current_idx}")
                    
                    # 3. Process the Image Dynamically based on current slider values
                    if orig_img.mode in ("RGBA", "P"): 
                        proc_img = orig_img.convert("RGB")
                    else:
                        proc_img = orig_img.copy()
                        
                    if proc_img.width > dyn_width:
                        ratio = dyn_width / float(proc_img.width)
                        new_h = int((float(proc_img.height) * float(ratio)))
                        proc_img = proc_img.resize((dyn_width, new_h), Image.Resampling.LANCZOS)
                        
                    img_byte_arr = io.BytesIO()
                    proc_img.save(img_byte_arr, format='WEBP', quality=dyn_quality, optimize=True)
                    new_bytes = img_byte_arr.getvalue()
                    new_kb = len(new_bytes) / 1024
                    
                    # 4. Render the Before & After Columns
                    col_l, col_r = st.columns(2)
                    with col_l:
                        st.markdown(f"**Original File:** {orig_img.width}x{orig_img.height}px | **{orig_kb:.1f} KB**")
                        st.image(orig_img, use_container_width=True)
                        
                    with col_r:
                        savings = 100 - ((new_kb / orig_kb) * 100) if orig_kb > 0 else 0
                        color = "#38ef7d" if savings > 0 else "#ff4b4b"
                        st.markdown(f"**WebP Preview:** {proc_img.width}x{proc_img.height}px | **{new_kb:.1f} KB** (<span style='color:{color}; font-weight:bold;'>-{savings:.1f}%</span>)", unsafe_allow_html=True)
                        st.image(new_bytes, use_container_width=True)
                    
                    st.write("") # Spacer
                    
                    # 5. Action Buttons
                    c_btn1, c_btn2, c_btn3 = st.columns([1, 1, 2])
                    if c_btn1.button("✅ Accept & Next", type="primary", use_container_width=True):
                        st.session_state.int_processed_imgs[current_file.name] = new_bytes
                        st.session_state.int_img_index += 1
                        st.rerun()
                        
                    if c_btn2.button("⏭️ Skip Image", use_container_width=True):
                        st.session_state.int_img_index += 1
                        st.rerun()
                        
                except Exception as e:
                    st.error(f"Error loading image: {e}")
                    if st.button("⏭️ Skip Broken File"):
                        st.session_state.int_img_index += 1
                        st.rerun()
                        
            else:
                # 6. Finish Screen
                st.success(f"🎉 All {len(up_imgs)} images reviewed!")
                
                if st.session_state.int_processed_imgs:
                    st.write(f"You optimized **{len(st.session_state.int_processed_imgs)}** images.")
                    zip_buf = io.BytesIO()
                    with zipfile.ZipFile(zip_buf, "a", zipfile.ZIP_DEFLATED, False) as zf:
                        for name, _bytes in st.session_state.int_processed_imgs.items():
                            orig_name = name.rsplit('.', 1)[0]
                            zf.writestr(f"{orig_name}_optimized.webp", _bytes)
                    
                    st.download_button("📥 Download Final ZIP", zip_buf.getvalue(), "interactive_optimized_images.zip", "application/zip", type="primary")
                else:
                    st.warning("You skipped all images. Nothing to download.")
                
                if st.button("🔄 Start Over"):
                    st.session_state.int_img_index = 0
                    st.session_state.int_processed_imgs = {}
                    st.rerun()

        st.divider()

    # ---------------------------------------------------------
    # TAB 2: PDF TOOLS
    # ---------------------------------------------------------
    with tab_pdf:
        st.subheader("📸 PDF to Image Converter")
        st.info("Extracts every page of a PDF into high-quality JPGs or PNGs.")
        pdf_file = st.file_uploader("Upload PDF File", type=['pdf'], key="pdf_to_img")
        img_fmt = st.radio("Output Image Format:", ["png", "jpeg"], horizontal=True)
        
        if pdf_file and st.button("📸 Convert PDF to Images", type="primary"):
            with st.spinner("Extracting pages..."):
                try:
                    doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
                    zip_buf = io.BytesIO()
                    with zipfile.ZipFile(zip_buf, "a", zipfile.ZIP_DEFLATED) as zf:
                        for page_num in range(len(doc)):
                            page = doc.load_page(page_num)
                            pix = page.get_pixmap(dpi=300) # High quality
                            img_bytes = pix.tobytes(img_fmt)
                            
                            orig_name = pdf_file.name.replace('.pdf', '')
                            ext = "jpg" if img_fmt == "jpeg" else "png"
                            zf.writestr(f"{orig_name}_page_{page_num+1}.{ext}", img_bytes)
                            
                    st.success(f"✅ Extracted {len(doc)} pages!")
                    st.download_button("📥 Download Images (ZIP)", zip_buf.getvalue(), f"{orig_name}_images.zip", "application/zip")
                except Exception as e:
                    st.error(f"Failed: {e}")

        st.divider()

        st.subheader("📝 PDF to Word (DOCX) Converter")
        pdf_to_word = st.file_uploader("Upload PDF", type=['pdf'], key="pdf_to_word")
        if pdf_to_word and st.button("📝 Convert to Word", type="primary"):
            with st.spinner("Converting document..."):
                try:
                    # Write to temp file (pdf2docx requires file paths, not streams)
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_in:
                        tmp_in.write(pdf_to_word.read())
                        tmp_in_path = tmp_in.name
                        
                    out_path = tmp_in_path.replace(".pdf", ".docx")
                    
                    cv = Converter(tmp_in_path)
                    cv.convert(out_path)
                    cv.close()
                    
                    with open(out_path, "rb") as f:
                        docx_data = f.read()
                        
                    st.success("✅ Converted to Word!")
                    orig_name = pdf_to_word.name.replace('.pdf', '.docx')
                    st.download_button("📥 Download DOCX", docx_data, orig_name, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                    
                    # Cleanup temp files
                    os.remove(tmp_in_path)
                    os.remove(out_path)
                except Exception as e:
                    st.error(f"Failed: {e}")

        st.divider()
        st.subheader("📄 Word (DOCX) to PDF Converter")
        st.info("Note: This feature requires Microsoft Word to be installed on the machine running Streamlit.")
        word_to_pdf = st.file_uploader("Upload Word Doc", type=['docx'], key="word_to_pdf")
        if word_to_pdf and st.button("📄 Convert to PDF", type="primary"):
            with st.spinner("Converting to PDF..."):
                try:
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp_in:
                        tmp_in.write(word_to_pdf.read())
                        tmp_in_path = tmp_in.name
                        
                    out_path = tmp_in_path.replace(".docx", ".pdf")
                    docx2pdf_convert(tmp_in_path, out_path)
                    
                    with open(out_path, "rb") as f:
                        pdf_data = f.read()
                        
                    st.success("✅ Converted to PDF!")
                    orig_name = word_to_pdf.name.replace('.docx', '.pdf')
                    st.download_button("📥 Download PDF", pdf_data, orig_name, "application/pdf")
                    
                    os.remove(tmp_in_path)
                    os.remove(out_path)
                except Exception as e:
                    st.error(f"Failed: Ensure MS Word is installed. Error: {e}")

    # ---------------------------------------------------------
    # TAB 3: HTML & DOC TOOLS
    # ---------------------------------------------------------
    with tab_html:
        st.subheader("🌐 HTML to Word (DOCX)")
        st.info("Paste HTML code and instantly get a cleanly formatted Word Document.")
        
        html_input = st.text_area("Paste HTML Code Here", height=200)
        if st.button("📝 Generate Word Doc", type="primary"):
            if not html_input:
                st.warning("Please paste HTML first.")
            else:
                try:
                    document = docx.Document()
                    new_parser = HtmlToDocx()
                    new_parser.add_html_to_document(html_input, document)
                    
                    doc_io = io.BytesIO()
                    document.save(doc_io)
                    
                    st.success("✅ Generated successfully!")
                    st.download_button("📥 Download Word Doc", doc_io.getvalue(), "HTML_Conversion.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                except Exception as e:
                    st.error(f"Error parsing HTML: {e}")

        st.divider()

        st.subheader("🌐 HTML to PDF")
        st.info("⚠️ Note: Requires `wkhtmltopdf` installed on your server/computer for this to work.")
        html_pdf_input = st.text_area("Paste HTML Code Here", height=200, key="html_to_pdf")
        if st.button("📄 Generate PDF", type="primary"):
            if not html_pdf_input:
                st.warning("Please paste HTML first.")
            else:
                try:
                    # Uses default wkhtmltopdf installation path.
                    pdf_bytes = pdfkit.from_string(html_pdf_input, False)
                    st.success("✅ PDF Generated!")
                    st.download_button("📥 Download PDF", pdf_bytes, "HTML_Conversion.pdf", "application/pdf")
                except Exception as e:
                    st.error(f"Failed. Make sure wkhtmltopdf is installed. Error: {e}")
